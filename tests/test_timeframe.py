# -*- coding: utf-8 -*-
"""
تست‌های موتور زمانی. اجرا:  python3 tests/test_timeframe.py

فلسفه این تست‌ها: مهم‌ترین چیزی که باید اثبات شود، «پیدا کردن چرخه» نیست —
«پیدا نکردنِ چرخه وقتی چرخه‌ای وجود ندارد» است. یک آشکارساز چرخه که روی نویز
هم چرخه می‌بیند، بدتر از نداشتن آشکارساز است.
"""
import datetime
import math
import os
import random
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))

from timeframe import elliott, fib_time, filters, gann, hurst, scoring, series, spectral
from timeframe.config import SpectralConfig
from timeframe.jalali import to_gregorian_date, to_jalali_date
from timeframe.macro import gold, iran, long_waves

FAILS = []


def check(name, cond, detail=""):
    status = "✔" if cond else "✘"
    if not cond:
        FAILS.append(name)
    print("  %s %s%s" % (status, name, ("  — " + detail) if detail else ""))


def section(t):
    print("\n" + t)
    print("─" * 70)


# ------------------------------------------------------------------ 1
def test_spectral_recovery():
    section("۱) بازیابی چرخه معلوم")
    for period in (25, 40, 60, 120):
        ts = series.synthetic("s", 900, [(period, 10, 0.3)], noise=1.5, trend=0.04)
        c = spectral.dominant_cycle(ts.closes)
        ok = c is not None and abs(c.period - period) / period < 0.06
        check("چرخه %d بازیابی شد" % period, ok,
              "یافت: %s" % (round(c.period, 1) if c else "هیچ"))


def test_spectral_two_cycles():
    section("۲) دو چرخه همزمان")
    ts = series.synthetic("s", 1000, [(80, 12, 0.0), (25, 5, 1.2)], noise=1.5)
    found = [c.period for c in spectral.find_cycles(ts.closes, SpectralConfig(top_n=4))
             if c.significant]
    check("هر دو چرخه ۸۰ و ۲۵ یافت شدند",
          any(abs(p - 80) < 6 for p in found) and any(abs(p - 25) < 3 for p in found),
          "یافت: %s" % [round(p, 1) for p in found])


def test_noise_control():
    section("۳) کنترل منفی — نویز سفید نباید چرخه بدهد")
    bad = 0
    for seed in range(8):
        rng = random.Random(seed)
        bars = [series.Bar(datetime.date(2020, 1, 1) + datetime.timedelta(days=i),
                           100 + rng.gauss(0, 3)) for i in range(800)]
        if spectral.dominant_cycle(series.TimeSeries("n", bars).closes) is not None:
            bad += 1
    check("۸ سری نویز، صفر چرخه معنادار", bad == 0, "مثبت کاذب: %d" % bad)


def test_random_walk_control():
    section("۴) کنترل منفی — گشت تصادفی (خودهمبسته)")
    bad = 0
    for seed in range(8):
        rng = random.Random(500 + seed)
        p, bars = 100.0, []
        for i in range(800):
            p *= 1 + rng.gauss(0.0003, 0.02)
            bars.append(series.Bar(datetime.date(2020, 1, 1) + datetime.timedelta(days=i), p))
        if spectral.dominant_cycle(series.TimeSeries("rw", bars).closes) is not None:
            bad += 1
    check("۸ گشت تصادفی، صفر چرخه معنادار", bad == 0, "مثبت کاذب: %d" % bad)


def test_goertzel_matches_dft():
    section("۵) صحت Goertzel در برابر DFT مستقیم")
    rng = random.Random(1)
    y = [math.sin(2 * math.pi * i / 32) * 5 + rng.gauss(0, 0.5) for i in range(256)]
    amp, _ph, re, im = spectral.goertzel(y, 32)
    n = len(y)
    re2 = sum(v * math.cos(2 * math.pi * i / 32) for i, v in enumerate(y))
    im2 = -sum(v * math.sin(2 * math.pi * i / 32) for i, v in enumerate(y))
    amp2 = 2 * math.hypot(re2, im2) / n
    check("دامنه Goertzel با DFT مستقیم یکی است", abs(amp - amp2) < 1e-9,
          "%.10f در برابر %.10f" % (amp, amp2))


def test_filters():
    section("۶) فیلتر میان‌گذر و تشخیص کف")
    ts = series.synthetic("s", 600, [(40, 10, 0.0)], noise=0.5, trend=0.05)
    tr = filters.find_extrema(filters.bandpass(ts.closes, 40), 24, "trough")
    gaps = [tr[i + 1] - tr[i] for i in range(len(tr) - 1)]
    check("فاصله کف‌ها ≈ ۴۰", gaps and all(abs(g - 40) <= 2 for g in gaps),
          "فواصل: %s" % gaps[:6])
    check("حذف روند خطی", abs(sum(filters.linear_detrend([i * 2.0 for i in range(50)]))) < 1e-9)


def test_hurst():
    section("۷) هرست: دوره مشاهده‌شده، لبه، FLD")
    ts = series.synthetic("s", 900, [(60, 12, 0.0)], noise=1.5, trend=0.03)
    st = hurst.analyze_cycle(ts, 60)
    check("دوره مشاهده‌شده نزدیک ۶۰", abs(st.observed_period - 60) < 2,
          "%.2f" % st.observed_period)
    check("فاز در بازه ۰ تا ۱", 0 <= st.phase < 1, "%.3f" % st.phase)
    check("ناحیه کور گزارش شده و منطقی", 0 < st.edge_blind_bars < len(ts) / 2,
          "%d کندل" % st.edge_blind_bars)
    f = hurst.fld_signal(ts, 60)
    check("FLD محاسبه شد", f.above is not None)
    shift = 30
    check("FLD همان قیمت جابه‌جاشده است",
          abs(hurst.fld(ts.closes, 60)[-1] - ts.closes[-1 - shift]) < 1e-9)


def test_gann():
    section("۸) گن")
    anchor = datetime.date(2026, 5, 30)
    w = gann.time_counts(anchor, "کف")
    days = {(x.date - anchor).days for x in w if x.kind == "day_count"}
    check("شمارش‌های روزانه درست‌اند", {30, 45, 60, 90, 144, 180} <= days)
    ann = [x for x in w if x.kind == "anniversary" and "1 " in x.label]
    check("سالگرد یک‌ساله درست است", ann and ann[0].date == datetime.date(2027, 5, 30))
    sq = gann.squaring(1000, 2000, 1000)
    check("مربع‌شدن مقیاس ۱ را می‌یابد", abs(sq.natural_scale - 1.0) < 1e-9
          and sq.nearest_gann_scale == 1)
    act = gann.active_windows([(datetime.date(2026, 8, 6), "کف")], datetime.date(2026, 9, 5))
    check("پنجره ۳۰ روزه فعال تشخیص داده شد", any("30 روز" in a.label for a in act),
          "%d پنجره فعال" % len(act))


def test_fib():
    section("۹) فیبوناچی زمانی")
    sw = fib_time.Swing(datetime.date(2026, 1, 1), datetime.date(2026, 4, 11), "س")
    ws = fib_time.project(sw)
    d618 = [w for w in ws if "0.618" in w.label][0]
    check("پروجکشن ۰٫۶۱۸ درست است",
          d618.date == datetime.date(2026, 4, 11) + datetime.timedelta(days=round(100 * 0.618)),
          str(d618.date))


def test_elliott():
    section("۱۰) قواعد الیوت")
    P, D = elliott.Pivot, datetime.date
    good = [P("0", D(2025, 1, 10), 1000), P("1", D(2025, 3, 1), 1300),
            P("2", D(2025, 4, 1), 1115), P("3", D(2025, 8, 1), 1800),
            P("4", D(2025, 9, 15), 1540), P("5", D(2026, 1, 20), 2100)]
    check("ایمپالس معتبر پذیرفته شد", elliott.validate(good).valid)
    r2 = list(good); r2[2] = P("2", D(2025, 4, 1), 950)
    check("نقض قاعده ۱ گرفته شد", "قاعده ۱" in " ".join(elliott.validate(r2).violations))
    r3 = list(good); r3[3] = P("3", D(2025, 8, 1), 1380); r3[4] = P("4", D(2025, 9, 15), 1330)
    check("نقض قاعده ۲ گرفته شد", "قاعده ۲" in " ".join(elliott.validate(r3).violations))
    r4 = list(good); r4[4] = P("4", D(2025, 9, 15), 1250)
    check("نقض قاعده ۳ گرفته شد", "قاعده ۳" in " ".join(elliott.validate(r4).violations))
    abc = [P("0", D(2026, 1, 20), 2100), P("A", D(2026, 3, 1), 1700),
           P("B", D(2026, 4, 10), 1900), P("C", D(2026, 6, 15), 1480)]
    check("ABC معتبر پذیرفته شد", elliott.validate(abc).valid)


def test_jalali():
    section("۱۱) تقویم شمسی")
    check("۱ فروردین ۱۴۰۳", to_jalali_date(datetime.date(2024, 3, 20)) == (1403, 1, 1))
    d, bad = datetime.date(2015, 1, 1), 0
    for _ in range(4000):
        if to_gregorian_date(*to_jalali_date(d)) != d:
            bad += 1
        d += datetime.timedelta(days=1)
    check("رفت‌وبرگشت ۴۰۰۰ روزه بدون خطا", bad == 0, "خطا: %d" % bad)


def test_gold():
    section("۱۲) طلا")
    b = gold.breakdown(3400, 950_000, 900_000_000)
    manual = 8.133 * 0.9 * (3400 / 31.1034768) * 950_000
    check("ارزش ذاتی با حساب دستی یکی است", abs(b.intrinsic_rial - manual) < 1)
    check("حباب درست محاسبه شد", abs(b.bubble_pct - (900_000_000 / manual - 1)) < 1e-12)
    d = gold.decompose_return(2600, 3400, 600_000, 950_000, 500_000_000, 900_000_000)
    check("تفکیک بازده به کل بازمی‌گردد",
          abs(d["کنترل ترکیبی"] - d["بازده کل ریالی"]) < 1e-9)


def test_macro():
    section("۱۳) رجیستری کلان")
    snap = long_waves.snapshot(datetime.date(2026, 9, 5))
    check("چرخه‌های تاریخ‌دار فاز جاری دارند", len(snap) >= 4, "%d چرخه" % len(snap))
    howe = [s for s in snap if s["key"] == "howe"][0]
    check("هاو در چرخش چهارم است", "بحران" in howe["phase"], howe["phase"])
    perez = [s for s in snap if s["key"] == "perez"][0]
    check("پرز در فاز استقرار موج ۵ است", "موج ۵" in perez["phase"], perez["phase"])
    check("هر چرخه منبع دارد", all(s["source"] for s in snap))
    check("هر چرخه سطح اطمینان تاریخ دارد", all(s["dating"] for s in snap))
    check("چارچوب‌های بدون تاریخ جدا شده‌اند", len(long_waves.undated_frameworks()) >= 5)


def test_iran_calendar():
    section("۱۴) تقویم ایران")
    c = iran.calendar_context(datetime.date(2027, 1, 20))
    titles = {w["title"] for w in c["windows"] if w["active"]}
    check("ناترازی گاز در دی فعال است", "ناترازی گاز زمستان" in titles, str(titles))
    c2 = iran.calendar_context(datetime.date(2026, 8, 20))
    check("ناترازی گاز در مرداد فعال نیست",
          "ناترازی گاز زمستان" not in {w["title"] for w in c2["windows"] if w["active"]})
    check("فصل مالی درست است", iran.calendar_context(datetime.date(2026, 9, 5))["fiscal_quarter"] == 2)


def test_scoring():
    section("۱۵) امتیازدهی")
    ts = series.synthetic("s", 900, [(60, 12, 0.0)], noise=1.5)
    st = hurst.analyze_cycle(ts, 60)
    sc = scoring.compute(cycle=st, fld=hurst.fld_signal(ts, 60), cycle_validated=True)
    check("امتیاز در بازه مجاز است", -10 <= sc.total <= 10, "%.2f" % sc.total)
    empty = scoring.compute()
    check("نبود اجزا هشدار می‌دهد و صفر می‌شود",
          empty.total == 0 and len(empty.warnings) >= 1)
    unval = scoring.compute(cycle=st, fld=hurst.fld_signal(ts, 60), cycle_validated=False)
    check("چرخه تأییدنشده هشدار می‌دهد",
          any("آزمون معناداری" in w for w in unval.warnings))
    check("چرخه تأییدنشده اتکاپذیری را پایین می‌آورد",
          unval.reliability != "بالا", unval.reliability)




# ------------------------------------------------------------------ بک‌تست و پیش‌بینی
def test_backtest_engine():
    section("۱۶) موتور بک‌تست — تفکیک لبه واقعی از شانس")
    from timeframe.backtest import Backtester, Costs, summarize
    from timeframe.backtest.benchmark import random_timing_null

    def rising(ts, period=60):
        return [1.0 if math.cos(2 * math.pi * i / period) > 0 else -1.0
                for i in range(len(ts))]

    bt = Backtester(costs=Costs())
    cyc = series.synthetic("c", 1200, [(60, 14, 0.0)], noise=1.0, trend=0.0, base=1000)
    r = bt.run(cyc, rising(cyc))
    p = summarize(r.returns, r.equity, r.positions, r.trade_pnls)
    check("روی چرخه واقعی سود می‌گیرد", p.total_return > 0.10,
          "%+.1f%%" % (p.total_return * 100))
    nt = random_timing_null(cyc, r, bt, n_samples=200)
    check("از زمان‌بندی تصادفی بهتر است", nt.p_value <= 0.05, "p=%.4f" % nt.p_value)

    bad = 0
    for seed in range(5):
        rng = random.Random(700 + seed)
        px, bars = 1000.0, []
        for i in range(1200):
            px *= 1 + rng.gauss(0.0002, 0.018)
            bars.append(series.Bar(datetime.date(2020, 1, 1) + datetime.timedelta(days=i),
                                   px, px * 1.01, px * 0.99))
        rw = series.TimeSeries("rw", bars)
        rr = bt.run(rw, rising(rw))
        n2 = random_timing_null(rw, rr, bt, n_samples=150)
        if n2.p_value <= 0.05:
            bad += 1
    check("روی ۵ گشت تصادفی لبه پیدا نمی‌کند", bad == 0, "مثبت کاذب: %d" % bad)


def test_backtest_no_lookahead():
    section("۱۷) نبود نگاه به آینده و اثر هزینه")
    from timeframe.backtest import Backtester, Costs

    def rising(ts, period=60):
        return [1.0 if math.cos(2 * math.pi * i / period) > 0 else -1.0
                for i in range(len(ts))]

    cyc = series.synthetic("c", 1200, [(60, 14, 0.0)], noise=1.0, base=1000)
    outs = []
    for lag in (1, 2, 5, 10):
        r = Backtester(costs=Costs(), execution_lag=lag).run(cyc, rising(cyc))
        outs.append(r.equity[-1] - 1.0)
    check("تأخیر بیشتر اجرا، بازده کمتر", outs[0] > outs[-1],
          "تأخیر۱ %+.1f%% ← تأخیر۱۰ %+.1f%%" % (outs[0] * 100, outs[-1] * 100))
    costs_out = []
    for slip in (0, 15, 50, 150):
        r = Backtester(costs=Costs(slippage_bps=slip)).run(cyc, rising(cyc))
        costs_out.append(r.equity[-1] - 1.0)
    check("هزینه بیشتر، بازده کمتر",
          all(costs_out[i] > costs_out[i + 1] for i in range(len(costs_out) - 1)),
          " → ".join("%+.0f%%" % (v * 100) for v in costs_out))


def test_deflated_sharpe():
    section("۱۸) شارپ تعدیل‌شده و توابع آماری")
    from timeframe.backtest import metrics as M
    for pr, z in ((0.975, 1.959964), (0.995, 2.575829), (0.5, 0.0)):
        check("norm_ppf(%.3f)" % pr, abs(M.norm_ppf(pr) - z) < 2e-4)
    check("norm_cdf(1.96) = 0.975", abs(M.norm_cdf(1.959964) - 0.975) < 1e-5)
    vals = [M.deflated_sharpe(0.08, 500, -0.3, 6.0, n) for n in (1, 10, 100, 1000)]
    check("تلاش بیشتر، شارپ تعدیل‌شده کمتر",
          all(vals[i] > vals[i + 1] for i in range(len(vals) - 1)),
          " → ".join("%.3f" % v for v in vals))


def test_regime_model():
    section("۱۹) مدل رژیم مارکوف — بازیابی پارامتر")
    from timeframe.forecast import regime as RG
    rng = random.Random(5)
    P = [[0.98, 0.02], [0.06, 0.94]]
    mu = [0.0012, -0.0025]
    sd = [0.010, 0.028]
    st, rets, truth = 0, [], []
    for _ in range(2500):
        rets.append(rng.gauss(mu[st], sd[st]))
        truth.append(st)
        st = 0 if rng.random() < P[st][0] else 1
    m = RG.fit(rets)
    s, c = m.stress_state, m.calm_state
    check("میانگین رژیم آرام بازیابی شد", abs(m.mu[c] - mu[0]) < 0.0005,
          "%+.5f در برابر %+.5f" % (m.mu[c], mu[0]))
    check("نوسان رژیم پرتنش بازیابی شد", abs(m.sigma[s] - sd[1]) < 0.005,
          "%.5f در برابر %.5f" % (m.sigma[s], sd[1]))
    acc = sum(1 for t in range(len(rets)) if (m.smoothed[t][s] > 0.5) == (truth[t] == 1))
    check("دقت تشخیص رژیم بالای ۸۵٪", acc / len(rets) > 0.85,
          "%.1f%%" % (100 * acc / len(rets)))
    check("احتمال‌ها در بازه مجاز", 0 <= m.prob_enter_stress(63) <= 1)
    sim = RG.simulate(m, 63, n_paths=800)
    check("مونت‌کارلو احتمال معتبر می‌دهد",
          0 <= sim["p_dd_20"] <= 1 and sim["n_paths"] == 800)


def test_base_rates():
    section("۲۰) نرخ پایه و اندازه نمونه مؤثر")
    from timeframe.forecast import base_rates as BR, drawdown as DD
    rng = random.Random(3)
    px, bars = 1000.0, []
    for i in range(2500):
        px *= 1 + rng.gauss(0.0004, 0.017)
        bars.append(series.Bar(datetime.date(2016, 1, 1) + datetime.timedelta(days=i),
                               px, px * 1.012, px * 0.988))
    ts = series.TimeSeries("t", bars)
    d = BR.unconditional(ts, 63)
    check("اندازه نمونه مؤثر ≈ خام ÷ افق",
          abs(d.n_effective - d.n_raw / 63) < 1e-6,
          "خام %d ← مؤثر %.0f" % (d.n_raw, d.n_effective))
    check("احتمال‌ها در بازه مجاز",
          all(0 <= x <= 1 for x in (d.p_negative, d.p_below_10, d.p_below_20)))
    fr = BR.forward_returns(ts, 10)
    manual = ts.closes[10] / ts.closes[0] - 1.0
    check("بازده آتی درست محاسبه شده", abs(fr[0] - manual) < 1e-12)
    check("انتهای سری بازده آتی ندارد", fr[-1] is None)
    dd = DD.forward_max_drawdown(ts, 63)
    check("افت‌ها منفی یا صفرند", dd.median_dd <= 0 and dd.worst <= dd.q95_dd)


def test_brier():
    section("۲۱) امتیاز بریر و تجزیه مرفی")
    from timeframe.forecast.scoring import brier
    check("پیش‌بین کامل بریر صفر می‌گیرد",
          abs(brier([1.0, 0.0, 1.0], [1, 0, 1]).brier) < 1e-12)
    check("پیش‌بین کاملاً غلط بریر ۱ می‌گیرد",
          abs(brier([0.0, 1.0], [1, 0]).brier - 1.0) < 1e-12)
    check("گفتن همیشگی ۵۰٪ بریر ۰٫۲۵ می‌دهد",
          abs(brier([0.5] * 10, [1, 0] * 5).brier - 0.25) < 1e-12)
    rng = random.Random(1)
    cf, co, of, oo = [], [], [], []
    for _ in range(300):
        pt = rng.random()
        o = 1 if rng.random() < pt else 0
        cf.append(pt); co.append(o)
        of.append(min(0.99, max(0.01, pt * 1.6 - 0.3))); oo.append(o)
    bc, bo = brier(cf, co), brier(of, oo)
    check("کالیبره از بیش‌اعتماد بهتر تشخیص داده می‌شود",
          bc.reliability < bo.reliability,
          "کالیبراسیون %.4f در برابر %.4f" % (bc.reliability, bo.reliability))
    check("اتحاد مرفی روی بریر سطل‌بندی‌شده دقیق است",
          abs((bc.reliability - bc.resolution + bc.uncertainty) - bc.brier_binned) < 1e-9,
          "باقی‌مانده %.2e" % abs((bc.reliability - bc.resolution + bc.uncertainty)
                                  - bc.brier_binned))
    disc_f = [round(x * 10) / 10 for x in cf]
    bd = brier(disc_f, co, n_bins=11)
    check("با پیش‌بینی گسسته، اتحاد روی بریر خام هم برقرار است",
          abs((bd.reliability - bd.resolution + bd.uncertainty) - bd.brier) < 1e-9,
          "باقی‌مانده %.2e" % abs((bd.reliability - bd.resolution + bd.uncertainty)
                                  - bd.brier))


def test_judgment_register():
    section("۲۲) دفتر پیش‌بینی")
    import tempfile
    from timeframe.forecast.judgment import Register, seed_templates
    path = os.path.join(tempfile.mkdtemp(), "fc.json")
    reg = Register(path)
    check("سؤال بدون معیار حل رد می‌شود",
          _raises(lambda: reg.add("X", "بازار بد می‌شود", "", "2026-12-01")))
    added = seed_templates(reg, "2026-12-05")
    check("الگوهای آماده افزوده شدند", len(added) >= 5)
    reg.forecast("MKT-DRAWDOWN-3M", 0.25, "نرخ پایه")
    reg.forecast("MKT-DRAWDOWN-3M", 0.35, "به‌روزرسانی")
    q = reg.questions["MKT-DRAWDOWN-3M"]
    check("به‌روزرسانی ثبت شد", len(q.forecasts) == 2 and abs(q.drift - 0.10) < 1e-9)
    reg.forecast("MKT-REGIME-3M", 1.0)
    check("قطعیت مطلق به ۰٫۹۹۵ محدود می‌شود",
          abs(reg.questions["MKT-REGIME-3M"].current - 0.995) < 1e-9)
    reg.resolve("MKT-DRAWDOWN-3M", 1)
    check("سؤال حل‌شده پیش‌بینی جدید نمی‌پذیرد",
          _raises(lambda: reg.forecast("MKT-DRAWDOWN-3M", 0.5)))
    reg.save()
    reg2 = Register(path)
    check("ذخیره و بارگذاری سالم است",
          len(reg2.questions) == len(reg.questions)
          and reg2.questions["MKT-DRAWDOWN-3M"].resolved)


def _raises(fn):
    try:
        fn()
        return False
    except Exception:
        return True


def test_outlook():
    section("۲۳) چشم‌انداز احتمالاتی")
    from timeframe.forecast import outlook as OL
    rng = random.Random(11)
    px, bars, reg = 1_000_000.0, [], 0
    for i in range(2000):
        if rng.random() < (0.012 if reg == 0 else 0.05):
            reg = 1 - reg
        px *= 1 + rng.gauss(0.0012 if reg == 0 else -0.0022,
                            0.010 if reg == 0 else 0.026)
        bars.append(series.Bar(datetime.date(2018, 1, 1) + datetime.timedelta(days=i),
                               px, px * 1.012, px * 0.988))
    o = OL.build(series.TimeSeries("idx", bars), 63)
    check("مدل رژیم برازش شد", o.regime is not None and bool(o.regime.mu))
    check("مونت‌کارلو اجرا شد", o.sim is not None)
    check("همه احتمال‌ها در بازه مجازند",
          all(0 <= v <= 1 for v in (o.p_stress_now, o.p_enter_stress,
                                    o.sim["p_below_10"], o.base.p_negative)))
    check("گزارش تولید می‌شود", "جمع‌بندی" in o.report())
    short = OL.build(series.TimeSeries("s", bars[:100]), 63)
    check("داده کم هشدار می‌دهد", bool(short.warnings))


# ------------------------------------------------------------------ فایل‌های تفکیک‌شده
def test_workbook_split():
    section("۲۴) تفکیک فایل‌های اکسل")
    import tempfile
    from openpyxl import load_workbook
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))
    from workbooks import BUILDERS, build

    out = tempfile.mkdtemp()
    paths = {}
    for key in BUILDERS:
        paths[key] = build(key, out)
    check("هر پنج فایل ساخته شد", len(paths) == 5, "، ".join(BUILDERS))

    expected = {
        "stocks": {"Dashboard", "Signals", "Data_Input", "Calculations",
                   "Daily_History", "Market_Index", "Time_Link", "Watchlist",
                   "Settings", "API_Map", "Documentation"},
        "options": {"Options", "Underlying", "Settings", "API_Map", "Documentation"},
        "time": {"Time_Cycles", "Macro_Cycles", "Forecast", "Sources",
                 "Settings", "Documentation"},
        "gold": {"Gold_Dashboard", "Asset_Signals", "Coin_Bubble", "Gold_Input",
                 "Gold_History", "Hist_Gram18", "Hist_Ons",
                 "Settings", "Documentation"},
        "fx": {"FX_Dashboard", "Asset_Signals", "Spreads", "FX_Input",
               "FX_History", "Hist_USDT", "Hist_Nima",
               "Settings", "Documentation"},
    }
    for key, want in expected.items():
        got = set(load_workbook(paths[key]).sheetnames)
        check("شیت‌های %s درست‌اند" % key, got == want,
              "اضافه: %s | کم: %s" % (got - want, want - got) if got != want else "")

    # هیچ فرمولی نباید به شیتی ارجاع دهد که در همان فایل نیست
    import re
    ref_re = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)!\$?[A-Z]")
    for key, path in paths.items():
        wb = load_workbook(path)
        names = set(wb.sheetnames)
        bad = []
        for ws in wb:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for m in ref_re.findall(c.value):
                            if m not in names and m not in ("IF", "AND", "OR"):
                                bad.append((ws.title, c.coordinate, m))
        check("%s ارجاع به شیت غایب ندارد" % key, not bad,
              str(bad[:3]) if bad else "")

    # ارجاع بین‌فایلی نباید وجود داشته باشد
    for key, path in paths.items():
        wb = load_workbook(path)
        ext = []
        for ws in wb:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and "[" in c.value and "]" in c.value \
                            and c.value.startswith("="):
                        ext.append((ws.title, c.coordinate))
        check("%s ارجاع بین‌فایلی ندارد" % key, not ext, str(ext[:3]) if ext else "")

    sizes = {k: os.path.getsize(v) for k, v in paths.items()}
    # آستانه‌ها یکسان نیستند و نباید باشند: options و time هرکدام یک سری دارند،
    # ولی gold و fx هرکدام **سه** شیت تاریخچه دارند (سکه/گرم/اونس و
    # دلار/تتر/نیمایی) چون برای هر دارایی یک سیگنال مستقل ساخته می‌شود. سه
    # برابر شدن حجم، هزینه واقعی همان قابلیت است نه هدررفت.
    check("options و time زیر ۳۰۰ کیلوبایت‌اند",
          all(sizes[k] < 300_000 for k in ("options", "time")),
          "، ".join("%s %dKB" % (k, sizes[k] // 1024) for k in ("options", "time")))
    check("gold و fx زیر ۷۰۰ کیلوبایت‌اند",
          all(sizes[k] < 700_000 for k in ("gold", "fx")),
          "، ".join("%s %dKB" % (k, sizes[k] // 1024) for k in ("gold", "fx")))
    check("هیچ فایلی به اندازه فایل یکپارچه قبلی نیست (زیر ۱ مگابایت)",
          all(v < 1_000_000 for v in sizes.values()),
          "، ".join("%s %dKB" % (k, v // 1024) for k, v in sorted(sizes.items())))


def test_bridge_scripts():
    section("۲۵) پل داده‌ای بین فایل‌ها")
    import tempfile
    from openpyxl import load_workbook
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
    from workbooks import build
    import link_workbooks as LW

    out = tempfile.mkdtemp()
    st = build("stocks", out)
    op = build("options", out)

    # بدون بازمحاسبه، مقدار فرمول‌ها None است — پل باید بدون خطا رد شود
    n = LW.link_options(st, op)
    check("پل بدون بازمحاسبه هم بدون استثنا اجرا می‌شود", n >= 0, "%d ردیف" % n)

    # پانویس زیر جدول نباید به‌عنوان نماد خوانده شود
    ws = load_workbook(op)["Underlying"]
    syms = [ws.cell(row=r, column=1).value
            for r in range(5, 5 + LW.UNDER_ROWS)
            if ws.cell(row=r, column=1).value]
    check("هیچ نمادی متن طولانی نیست (پانویس خوانده نشده)",
          all(len(str(s)) < 20 for s in syms), str(syms))


def test_market_data_providers():
    section("۲۶) منابع داده طلا و ارز")
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
    from market_data.base import Quote, cross_check
    from market_data.providers import build_registry

    reg = build_registry()
    check("همه کمیت‌ها دست‌کم یک منبع دارند",
          all(reg.for_quantity(q) for q in
              ("gold_oz_usd", "usd_irr_free", "usdt_irr", "coin_full_irr")))
    check("منبع دستی همه کمیت‌ها را پوشش می‌دهد",
          any(p.key == "manual" and len(p.supplies) >= 13 for p in reg.providers))

    # آزمون سلامت مطلق
    check("دلار سالم پذیرفته می‌شود", Quote("usd_irr_free", 950_000, "t").ok)
    check("دلار به تومان رد می‌شود", not Quote("usd_irr_free", 95_000, "t").ok)
    check("اونس سالم پذیرفته می‌شود", Quote("gold_oz_usd", 3400, "t").ok)
    check("اونس با اسکیل غلط رد می‌شود", not Quote("gold_oz_usd", 3_400_000, "t").ok)

    # آزمون سازگاری متقابل — مستقل از سطح تورم
    def mk(d):
        return {k: Quote(k, v, "t") for k, v in d.items()}

    healthy = mk({"gold_oz_usd": 3412.5, "usd_irr_free": 952_000,
                  "usd_irr_nima": 718_000, "usdt_irr": 969_500,
                  "coin_full_irr": 905_000_000, "coin_half_irr": 472_000_000,
                  "coin_quarter_irr": 286_000_000, "gram18k_irr": 82_400_000,
                  "eur_irr": 1_035_000, "aed_irr": 259_300})
    check("داده سالم هیچ ناسازگاری نمی‌دهد", not cross_check(healthy),
          str(cross_check(healthy))[:120])

    ten_x = dict(healthy)
    ten_x["coin_full_irr"] = Quote("coin_full_irr", 90_500_000, "t")
    check("سکه با واحد غلط گرفته می‌شود", len(cross_check(ten_x)) >= 2,
          "%d ناسازگاری" % len(cross_check(ten_x)))

    swapped = dict(healthy)
    swapped["usd_irr_free"] = Quote("usd_irr_free", 718_000, "t")
    swapped["usd_irr_nima"] = Quote("usd_irr_nima", 952_000, "t")
    check("جابه‌جایی آزاد و نیمایی گرفته می‌شود",
          any("برعکس" in w for w in cross_check(swapped)))

    # تورم ۵ برابری نباید هشدار کاذب بدهد — آزمون نسبتی است، نه مطلق
    inflated = mk({"gold_oz_usd": 3412.5, "usd_irr_free": 4_760_000,
                   "usdt_irr": 4_850_000, "coin_full_irr": 4_525_000_000,
                   "coin_half_irr": 2_360_000_000,
                   "coin_quarter_irr": 1_430_000_000})
    check("تورم ۵ برابری هشدار کاذب نمی‌دهد", not cross_check(inflated),
          str(cross_check(inflated))[:120])


def test_manual_provider_roundtrip():
    section("۲۷) مسیر دستی طلا و ارز")
    import json
    import tempfile
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
    from market_data.providers import build_registry

    d = tempfile.mkdtemp()
    path = os.path.join(d, "m.json")
    payload = {"gold_oz_usd": 3400, "usd_irr_free": 950_000, "usdt_irr": 968_000,
               "coin_full_irr": 900_000_000}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh)

    reg = build_registry()
    q = reg.resolve({"manual_file": path, "timeout": 3}, order=["manual"])
    check("همه مقادیر دستی خوانده شدند", len(q) >= 4, "%d کمیت" % len(q))
    check("مقادیر درست‌اند",
          abs(q["usd_irr_free"].value - 950_000) < 1
          and abs(q["gold_oz_usd"].value - 3400) < 1)
    check("منشأ ثبت شده", all(x.source and x.fetched_at for x in q.values()))

    empty = os.path.join(d, "empty.json")
    with open(empty, "w", encoding="utf-8") as fh:
        json.dump({"چیز_نامربوط": 1}, fh)
    from market_data.providers import manual as MP
    check("فایل بدون کلید معتبر خطای گویا می‌دهد",
          _raises(lambda: MP.fetch({"manual_file": empty})))
    check("فایل غایب خطای گویا می‌دهد",
          _raises(lambda: MP.fetch({"manual_file": os.path.join(d, "nope.json")})))


def main():
    for fn in (test_spectral_recovery, test_spectral_two_cycles, test_noise_control,
               test_random_walk_control, test_goertzel_matches_dft, test_filters,
               test_hurst, test_gann, test_fib, test_elliott, test_jalali,
               test_gold, test_macro, test_iran_calendar, test_scoring,
               test_backtest_engine, test_backtest_no_lookahead, test_deflated_sharpe,
               test_regime_model, test_base_rates, test_brier,
               test_judgment_register, test_outlook,
               test_workbook_split, test_bridge_scripts,
               test_market_data_providers, test_manual_provider_roundtrip,
               test_history_parsers, test_valuation_series,
               test_asset_signals_sheet):
        fn()
    print("\n" + "═" * 70)
    if FAILS:
        print("✘ %d تست ناموفق: %s" % (len(FAILS), "، ".join(FAILS)))
        return 1
    print("✔ همه تست‌ها موفق.")
    return 0


def test_history_parsers():
    section("۲۸) تجزیه تاریخچه tgju و Nobitex")
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
    from market_data.providers import nobitex, tgju

    # --- tgju: ردیف‌ها [باز، کمترین، بیشترین، بسته، تغییرHTML، درصدHTML، شمسی، میلادی]
    payload = {"data": [
        ["1,000,000", "990,000", "1,010,000", "1,005,000",
         "<span class='high'>5,000</span>", "<span>0.5%</span>",
         "1404/06/15", "2025/09/06"],
        ["۱٬۰۰۵٬۰۰۰", "۱٬۰۰۰٬۰۰۰", "۱٬۰۲۰٬۰۰۰", "۱٬۰۱۸٬۰۰۰",
         "<span>13,000</span>", "<span>1.3%</span>", "1404/06/16", "2025-09-07"],
        ["bad", "bad", "bad", "bad", "", "", "", ""],          # باید نادیده گرفته شود
    ]}
    orig = tgju.get_json
    tgju.get_json = lambda u, timeout=None, referer=None: payload
    try:
        rows = tgju.fetch_history("sekee")
    finally:
        tgju.get_json = orig
    check("ردیف خراب دور ریخته می‌شود", len(rows) == 2, "%d ردیف" % len(rows))
    check("تاریخ میلادی خوانده می‌شود (نه شمسی)",
          rows[0][0] == datetime.date(2025, 9, 6), str(rows[0][0]))
    check("هر دو قالب تاریخ پشتیبانی می‌شوند",
          rows[1][0] == datetime.date(2025, 9, 7), str(rows[1][0]))
    check("HTML از ستون‌ها پاک می‌شود و عدد سالم است", rows[0][4] == 1_005_000.0)
    check("ارقام فارسی خوانده می‌شوند", rows[1][4] == 1_018_000.0, str(rows[1][4]))
    check("ترتیب صعودی است", rows[0][0] < rows[1][0])

    # --- Nobitex UDF: مقادیر IRT به **تومان**‌اند، پس باید ×۱۰ شوند
    udf = {"s": "ok", "t": [1562095800, 1562182200],
           "o": [146272500, 150551000], "h": [152000000, 158000000],
           "l": [140062400, 150551000], "c": [151440200, 157000000],
           "v": [18.2, 9.8]}
    orig = nobitex.get_json
    nobitex.get_json = lambda u, timeout=None: udf
    try:
        bars = nobitex.fetch_history("USDTIRT")
    finally:
        nobitex.get_json = orig
    check("دو کندل تجزیه شد", len(bars) == 2)
    check("تومان به ریال تبدیل می‌شود (×۱۰)",
          bars[0]["close"] == 151440200 * 10, str(bars[0]["close"]))
    check("زمان epoch ثانیه‌ای به تاریخ تبدیل می‌شود",
          bars[0]["date"] == datetime.date(2019, 7, 2), str(bars[0]["date"]))

    nobitex.get_json = lambda u, timeout=None: {"s": "no_data"}
    try:
        failed = False
        try:
            nobitex.fetch_history("USDTIRT")
        except ValueError:
            failed = True
        check("پاسخ no_data استثنا می‌دهد (نه سری خالی بی‌صدا)", failed)
    finally:
        nobitex.get_json = orig


def test_valuation_series():
    section("۲۹) سنجه ارزش‌گذاری طلا و ارز")
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
    from market_data import history as H

    d = datetime.date(2026, 1, 1)
    ons, usd = 3200.0, 1_000_000.0
    intrinsic = 8.133 * 0.900 * (ons / 31.1034768) * usd
    closes = {"coin_full": {d: intrinsic * 1.18}, "ons": {d: ons},
              "usd_free": {d: usd}}
    v = H.build_valuation("coin_bubble", closes)
    check("حباب سکه ۱۸٪ دقیقاً بازتولید می‌شود",
          abs(v[d] - 0.18) < 1e-9, "%.6f" % v[d])

    parity = 0.750 * (ons / 31.1034768) * usd
    v = H.build_valuation("gram_premium",
                          {"gram18": {d: parity * 1.07}, "ons": {d: ons},
                           "usd_free": {d: usd}})
    check("پریمیوم گرم ۷٪ بازتولید می‌شود", abs(v[d] - 0.07) < 1e-9)

    v = H.build_valuation("usdt_premium",
                          {"usdt": {d: usd * 1.02}, "usd_free": {d: usd}})
    check("پریمیوم تتر ۲٪ بازتولید می‌شود", abs(v[d] - 0.02) < 1e-9)

    # تله ریال/تومان: سکه به تومان → حباب حدود ۸۸٪− که بیرون دامنه معقول است
    bad = {"coin_full": {d: intrinsic * 1.18 / 10}, "ons": {d: ons},
           "usd_free": {d: usd}}
    check("سکه با واحد غلط در سنجه نوشته نمی‌شود",
          not H.build_valuation("coin_bubble", bad))

    # هم‌ترازی تاریخ: روزی که یکی از نهاده‌ها ندارد، سنجه نمی‌گیرد
    d2 = datetime.date(2026, 1, 2)
    mixed = {"coin_full": {d: intrinsic * 1.18, d2: intrinsic * 1.18},
             "ons": {d: ons}, "usd_free": {d: usd, d2: usd}}
    got = H.build_valuation("coin_bubble", mixed)
    check("روز بدون نهاده کامل کنار گذاشته می‌شود",
          set(got) == {d}, str(sorted(got)))

    # نبودِ کامل یک سری → سنجه ساخته نمی‌شود، نه اینکه صفر شود
    check("نبودِ یک سری سنجه را صفر نمی‌کند",
          H.build_valuation("coin_bubble", {"coin_full": {d: 1.0}}) == {})

    check("هر شیت تاریخچه یک سری تعریف‌شده دارد",
          all(v[0] in H.SERIES for v in H.SHEETS.values()))
    check("هر سنجه ارجاع‌شده تعریف شده است",
          all(v[1] in H.VALUATIONS for v in H.SHEETS.values() if v[1]))


def test_asset_signals_sheet():
    section("۳۰) شیت سیگنال دارایی‌های تک‌سری")
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))
    from workbooks import asset_signals

    keys = [c[0] for c in asset_signals.COLS]
    for k in ("price", "val", "val_pct", "t_score", "m_score", "v_score",
              "p_score", "total", "signal"):
        check("ستون %s وجود دارد" % k, k in keys)

    tmpl = dict((c[0], c[4]) for c in asset_signals.COLS if c[4])

    # هیچ فرمولی نباید به مقایسه با "" تکیه کند: در LibreOffice شرط 0="" درست
    # است ولی در اکسل نادرست — همان تله‌ای که قبلاً ردیف‌های خالی را سیگنال‌دار
    # نشان می‌داد.
    offenders = [k for k, t in tmpl.items() if '}=""' in t]
    check("هیچ فرمولی به تله 0=\"\" تکیه نمی‌کند",
          not offenders, "، ".join(offenders))

    check("خواندن سنجه با ISBLANK محافظت شده",
          "ISBLANK" in tmpl["val"] and "ISBLANK" in tmpl["val_pct"])
    check("مخرج امتیاز کل پویاست (وزن ارزش‌گذاری شرطی است)",
          "IF(ISNUMBER($Q{r}),AW_VAL,0)" in tmpl["total"])
    check("امتیاز کل به AW_SUM ثابت تقسیم نمی‌شود",
          "AW_SUM" not in tmpl["total"])
    for nm in ("TH_SBUY", "TH_BUY", "TH_SELL", "TH_SSELL"):
        check("آستانه %s از Settings خوانده می‌شود" % nm, nm in tmpl["signal"])



if __name__ == "__main__":
    sys.exit(main())
