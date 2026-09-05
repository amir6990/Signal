# -*- coding: utf-8 -*-
"""
common.py — استایل‌ها، ثابت‌ها و ابزارهای مشترک ساخت فایل سیگنال بازار سرمایه ایران.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

# ---------- رنگ‌ها ----------
C_HEADER_BG   = "1F3864"   # سرمه‌ای تیره — سربرگ
C_SUBHDR_BG   = "D9E2F3"   # آبی روشن — ردیف فیلد API
C_SECTION_BG  = "2E5C8A"
C_INPUT_BG    = "FFF2CC"   # زرد کم‌رنگ — سلول ورودی کاربر
C_NOTE        = "808080"
C_BLUE_INPUT  = "0000FF"   # ورودی hardcode
C_GREEN_LINK  = "008000"   # لینک بین‌شیتی
C_BLACK       = "000000"

C_SBUY_BG, C_SBUY_FG = "00B050", "FFFFFF"   # خرید قوی
C_BUY_BG,  C_BUY_FG  = "C6EFCE", "006100"   # خرید
C_HOLD_BG, C_HOLD_FG = "D9D9D9", "3F3F3F"   # نگهداری
C_SELL_BG, C_SELL_FG = "FFC7CE", "9C0006"   # فروش
C_SSELL_BG, C_SSELL_FG = "C00000", "FFFFFF" # فروش قوی

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------- فرمت عدد ----------
F_INT   = '#,##0;(#,##0);-'
F_RIAL  = '#,##0;(#,##0);-'
F_PRICE = '#,##0;(#,##0);-'
F_PCT   = '0.0%;(0.0%);-'
F_PCT2  = '0.00%;(0.00%);-'
F_NUM1  = '0.0;(0.0);-'
F_NUM2  = '0.00;(0.00);-'
F_X     = '0.00"x"'
F_DATE  = 'yyyy-mm-dd'
F_MRIAL = '#,##0,,"M";(#,##0,,"M");-'      # میلیون ریال
F_BRIAL = '#,##0,,,"B";(#,##0,,,"B");-'    # میلیارد ریال


def hdr(ws, row, headers, api_row=None, start_col=1, height=30):
    """سربرگ دو ردیفی: ردیف بالا برچسب فارسی، ردیف پایین کلید فیلد API."""
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=C_HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height
    if api_row is not None:
        for i, a in enumerate(api_row):
            c = ws.cell(row=row + 1, column=start_col + i, value=a)
            c.font = Font(name=FONT, italic=True, size=8, color=C_NOTE)
            c.fill = PatternFill("solid", fgColor=C_SUBHDR_BG)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
        ws.row_dimensions[row + 1].height = 22


def title_block(ws, text, subtitle=None, ncols=8):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=C_HEADER_BG)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, size=9, italic=True, color="404040")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws.row_dimensions[2].height = 18


def widths(ws, spec):
    """spec: dict {'A': 12, ...} یا list از عرض‌ها از ستون A"""
    if isinstance(spec, dict):
        for k, v in spec.items():
            ws.column_dimensions[k].width = v
    else:
        for i, v in enumerate(spec):
            ws.column_dimensions[get_column_letter(i + 1)].width = v


def style_data(ws, r1, r2, c1, c2, fmt=None, color=C_BLACK, bold=False, size=9,
               align="center", border=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=size, bold=bold, color=color)
            cell.alignment = Alignment(horizontal=align, vertical="center")
            if fmt:
                cell.number_format = fmt
            if border:
                cell.border = BORDER


def note(ws, cell, text, color=C_NOTE, size=8, italic=True, bold=False):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=size, italic=italic, bold=bold, color=color)
    ws[cell].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)


# ---------- تبدیل تاریخ میلادی به شمسی (الگوریتم استاندارد، بدون کتابخانه) ----------
_G_DM = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]


def to_jalali(gy, gm, gd):
    gy2 = gy - 1600
    gm2 = gm - 1
    gd2 = gd - 1
    g_day_no = 365 * gy2 + (gy2 + 3) // 4 - (gy2 + 99) // 100 + (gy2 + 399) // 400
    g_day_no += _G_DM[gm2] + gd2
    if gm > 2 and ((gy % 4 == 0 and gy % 100 != 0) or (gy % 400 == 0)):
        g_day_no += 1
    j_day_no = g_day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053
    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461
    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365
    for i in range(11):
        md = 31 if i < 6 else 30
        if j_day_no < md:
            break
        j_day_no -= md
    else:
        i = 11
    jm = i + 1
    jd = j_day_no + 1
    return jy, jm, jd


def jalali_str(d):
    jy, jm, jd = to_jalali(d.year, d.month, d.day)
    return "%04d/%02d/%02d" % (jy, jm, jd)
