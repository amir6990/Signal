# -*- coding: utf-8 -*-
"""سازنده مشترک فایل‌های دارایی تک‌سری (طلا، ارز) — ساختار یکسان، محتوای متفاوت."""
import datetime

from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as K
from common import (BORDER, F_DATE, F_INT, F_NUM1, F_NUM2, F_PCT, F_PCT2, F_PRICE,
                    F_X, FONT, hdr, note, style_data, title_block, widths)

# تاریخچه روزانه: ستون‌های خام (A..F) + محاسباتی (G..)
HIST_RAW = [
    ("تاریخ میلادی", "-", 13, F_DATE),
    ("تاریخ شمسی", "-", 13, None),
    ("قیمت", "close", 16, F_PRICE),
    ("بیشترین", "high", 14, F_PRICE),
    ("کمترین", "low", 14, F_PRICE),
    ("حجم / یادداشت", "volume", 14, F_INT),
]

HIST_CALC = [
    # شمارنده باید هم ردیف هدر متنی بالای خودش را تحمل کند و هم ردیف‌های خالی
    # پایین جدول را. ISNUMBER هر دو حالت را می‌گیرد؛ مقایسه مستقیم با عدد نه.
    ("شمارنده", "n", 9, "0", '=IF($A{r}="","",IF(ISNUMBER($G{p}),$G{p}+1,1))', 0),
    ("بازده روزانه", "-", 11, F_PCT2,
     '=IF(OR(NOT(ISNUMBER($G{r})),$G{r}<2,NOT(ISNUMBER($C{p})),$C{p}=0),"",$C{r}/$C{p}-1)', 0),
    ("MA کوتاه", "MA_SHORT", 14, F_PRICE,
     '=IF(AND(ISNUMBER($G{r}),$G{r}>=MA_SHORT),AVERAGE(INDEX($C${lo}:$C{r},{k}-MA_SHORT):$C{r}),"")', 400),
    ("MA میان‌مدت", "MA_MED", 14, F_PRICE,
     '=IF(AND(ISNUMBER($G{r}),$G{r}>=MA_MED),AVERAGE(INDEX($C${lo}:$C{r},{k}-MA_MED):$C{r}),"")', 400),
    ("MA بلند", "MA_LONG2", 14, F_PRICE,
     '=IF(AND(ISNUMBER($G{r}),$G{r}>=MA_LONG2),AVERAGE(INDEX($C${lo}:$C{r},{k}-MA_LONG2):$C{r}),"")', 400),
    ("نوسان ۲۰ روزه (سالانه)", "-", 15, F_PCT,
     '=IF(AND(ISNUMBER($G{r}),$G{r}>=21),STDEV(INDEX($H${lo}:$H{r},{k}-20):$H{r})*SQRT(245),"")', 60),
    ("سقف ۶۰ روزه", "-", 13, F_PRICE,
     '=IF(AND(ISNUMBER($G{r}),$G{r}>=60),MAX(INDEX($D${lo}:$D{r},{k}-60):$D{r}),"")', 200),
    ("کف ۶۰ روزه", "-", 13, F_PRICE,
     '=IF(AND(ISNUMBER($G{r}),$G{r}>=60),MIN(INDEX($E${lo}:$E{r},{k}-60):$E{r}),"")', 200),
    ("افت از سقف", "-", 12, F_PCT,
     '=IF(OR(NOT(ISNUMBER($C{r})),NOT(ISNUMBER($M{r})),$M{r}=0),"",$C{r}/$M{r}-1)', 0),
]

# ستون‌های ارزش‌گذاری — P خام (نوشته پایتون)، Q صدک (فرمول اکسل).
# چرا این تقسیم کار: ساختن سری حباب نیازمند هم‌ترازی سه سری زمانی با تاریخ‌های
# ناهمسان است که در اکسل شکننده می‌شود؛ در پایتون بدیهی است. اما صدک باید در
# اکسل باشد تا با تغییر داده زنده به‌روز شود.
VAL_RAW = ("سنجه ارزش‌گذاری", "پایتون", 16, F_PCT)
VAL_CALC = ("صدک تاریخی سنجه", "PERCENTRANK", 15, F_PCT,
            '=IF(NOT(ISNUMBER($P{r})),"",'
            'IFERROR(PERCENTRANK($P${first}:$P${last},$P{r}),""))')

FIRST_ROW = 5


def hist_formula(tmpl, back, r):
    lo = max(FIRST_ROW, r - back)
    return tmpl.format(r=r, p=r - 1, lo=lo, k=r - lo + 2)


def build_history_sheet(wb, name, title, subtitle, rows, n_spare=700,
                        valuation_label=None):
    """شیت تاریخچه با ستون‌های خام + محاسباتی.

    rows: list[(date, close, high, low, vol)]
    valuation_label: اگر داده شود، ستون سنجه ارزش‌گذاری و صدک آن اضافه می‌شود
        (مثلاً «حباب سکه» یا «پریمیوم تتر»).
    """
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ncol = len(HIST_RAW) + len(HIST_CALC) + (2 if valuation_label else 0)
    title_block(ws, title, subtitle, ncol)
    w = [c[2] for c in HIST_RAW] + [c[2] for c in HIST_CALC]
    h = [c[0] for c in HIST_RAW] + [c[0] for c in HIST_CALC]
    a = [c[1] for c in HIST_RAW] + [c[1] for c in HIST_CALC]
    if valuation_label:
        w += [VAL_RAW[2], VAL_CALC[2]]
        h += [valuation_label, VAL_CALC[0]]
        a += [VAL_RAW[1], VAL_CALC[1]]
    widths(ws, w)
    hdr(ws, 3, h, a)

    from timeframe.jalali import jalali_str
    total = max(len(rows), 0) + n_spare
    for i in range(total):
        r = FIRST_ROW + i
        if i < len(rows):
            d, c, h, l, v = rows[i]
            ws.cell(row=r, column=1, value=d).number_format = F_DATE
            ws.cell(row=r, column=2, value=jalali_str(d))
            ws.cell(row=r, column=3, value=c).number_format = F_PRICE
            ws.cell(row=r, column=4, value=h).number_format = F_PRICE
            ws.cell(row=r, column=5, value=l).number_format = F_PRICE
            ws.cell(row=r, column=6, value=v).number_format = F_INT
        else:
            for cc in range(1, 7):
                cell = ws.cell(row=r, column=cc)
                cell.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
                cell.font = Font(name=FONT, size=8, color=K.C_BLUE_INPUT)
                cell.border = BORDER
            ws.cell(row=r, column=1).number_format = F_DATE
            ws.cell(row=r, column=3).number_format = F_PRICE
        for j, (_l, _d, _w, fmt, tmpl, back) in enumerate(HIST_CALC):
            c = ws.cell(row=r, column=len(HIST_RAW) + 1 + j,
                        value=hist_formula(tmpl, back, r))
            if fmt:
                c.number_format = fmt
            c.font = Font(name=FONT, size=8)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER
    last = FIRST_ROW + total - 1
    if valuation_label:
        vcol = len(HIST_RAW) + len(HIST_CALC) + 1
        for i in range(total):
            r = FIRST_ROW + i
            cv = ws.cell(row=r, column=vcol)
            cv.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
            cv.font = Font(name=FONT, size=8, color=K.C_GREEN_LINK)
            cv.number_format = VAL_RAW[3]
            cv.border = BORDER
            cq = ws.cell(row=r, column=vcol + 1,
                         value=VAL_CALC[4].format(r=r, first=FIRST_ROW, last=last))
            cq.number_format = VAL_CALC[3]
            cq.font = Font(name=FONT, size=8)
            cq.alignment = Alignment(horizontal="center", vertical="center")
            cq.border = BORDER
    style_data(ws, FIRST_ROW, FIRST_ROW + len(rows) - 1, 1, len(HIST_RAW), size=8)
    for r in range(FIRST_ROW, FIRST_ROW + len(rows)):
        for i, col in enumerate(HIST_RAW):
            if col[3]:
                ws.cell(row=r, column=i + 1).number_format = col[3]
    ws.auto_filter.ref = "A4:%s%d" % (get_column_letter(ncol), last)
    ws.freeze_panes = "C5"
    return ws, last


LAST_ROW_FORMULA = '=COUNT($A$5:$A$%d)+4'

# ⚠️ تا وقتی شیت تاریخچه خالی است، فرمول بالا عدد ۴ می‌دهد — و ردیف ۴ همان
# ردیف سرستون است. پس هر INDEX زیر با «اشاره‌گر < ۵» محافظت شده تا به جای
# خالی، متنِ سرستون در بلوک وضعیت روند ظاهر نشود.


def trend_block(ws, sheet, last_row, start_row, label, ncol=6):
    """بلوک وضعیت روند از شیت تاریخچه — با یافتن آخرین ردیف دارای داده."""
    rows = [
        ("آخرین ردیف داده", LAST_ROW_FORMULA % last_row, "0"),
        ("تاریخ آخرین داده", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$A$1:$A${n},$B${r0}),\"\"))", F_DATE),
        ("تاریخ شمسی", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$B$1:$B${n},$B${r0}),\"\"))", None),
        ("قیمت", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$C$1:$C${n},$B${r0}),\"\"))", F_PRICE),
        ("بازده روز", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$H$1:$H${n},$B${r0}),\"\"))", F_PCT2),
        ("MA کوتاه", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$I$1:$I${n},$B${r0}),\"\"))", F_PRICE),
        ("MA میان‌مدت", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$J$1:$J${n},$B${r0}),\"\"))", F_PRICE),
        ("MA بلند", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$K$1:$K${n},$B${r0}),\"\"))", F_PRICE),
        ("نوسان سالانه", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$L$1:$L${n},$B${r0}),\"\"))", F_PCT),
        ("سقف ۶۰ روزه", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$M$1:$M${n},$B${r0}),\"\"))", F_PRICE),
        ("کف ۶۰ روزه", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$N$1:$N${n},$B${r0}),\"\"))", F_PRICE),
        ("افت از سقف", "=IF($B${r0}<5,\"\",IFERROR(INDEX({s}!$O$1:$O${n},$B${r0}),\"\"))", F_PCT),
    ]
    r = start_row
    r0 = start_row
    for lbl, f, fmt in rows:
        ws.cell(row=r, column=1, value=lbl).font = Font(name=FONT, size=9)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
        c = ws.cell(row=r, column=2,
                    value=f.format(s=sheet, n=last_row, r0=r0) if "{" in f else f)
        c.font = Font(name=FONT, size=10, bold=True, color=K.C_GREEN_LINK)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        if fmt:
            c.number_format = fmt
        r += 1
    # نمره روند ساده
    ws.cell(row=r, column=1, value="امتیاز روند (−۱۰ تا +۱۰)").font = Font(
        name=FONT, bold=True, size=10)
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
    sc = ws.cell(row=r, column=2, value=(
        '=IFERROR(MAX(-10,MIN(10,'
        'IF($B${p}>$B${ms},2.5,-2.5)+IF($B${p}>$B${mm},3.5,-3.5)+IF($B${p}>$B${ml},4,-4)'
        ')),0)').format(p=r0 + 3, ms=r0 + 5, mm=r0 + 6, ml=r0 + 7))
    sc.font = Font(name=FONT, bold=True, size=12)
    sc.number_format = F_NUM1
    sc.alignment = Alignment(horizontal="center", vertical="center")
    sc.border = BORDER
    ws.conditional_formatting.add("B%d" % r, ColorScaleRule(
        start_type="num", start_value=-10, start_color=K.C_SSELL_BG,
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=10, end_color=K.C_SBUY_BG))
    lbl_row = r + 1
    ws.cell(row=lbl_row, column=1, value="وضعیت روند").font = Font(name=FONT, bold=True, size=10)
    ws.cell(row=lbl_row, column=1).alignment = Alignment(horizontal="right", vertical="center")
    lb = ws.cell(row=lbl_row, column=2, value=(
        '=IF($B${s}>=5,"صعودی قوی",IF($B${s}>=2,"صعودی",'
        'IF($B${s}<=-5,"نزولی قوی",IF($B${s}<=-2,"نزولی","خنثی"))))').format(s=r))
    lb.font = Font(name=FONT, bold=True, size=11)
    lb.alignment = Alignment(horizontal="center", vertical="center")
    lb.border = BORDER
    for lbl_txt, bg, fg in (("صعودی قوی", K.C_SBUY_BG, K.C_SBUY_FG),
                            ("صعودی", K.C_BUY_BG, K.C_BUY_FG),
                            ("خنثی", K.C_HOLD_BG, K.C_HOLD_FG),
                            ("نزولی", K.C_SELL_BG, K.C_SELL_FG),
                            ("نزولی قوی", K.C_SSELL_BG, K.C_SSELL_FG)):
        ws.conditional_formatting.add("B%d" % lbl_row, CellIsRule(
            operator="equal", formula=['"%s"' % lbl_txt],
            fill=PatternFill("solid", fgColor=bg), font=Font(color=fg, bold=True)))
    return lbl_row, r
