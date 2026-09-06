# -*- coding: utf-8 -*-
"""
شیت‌های تحلیل زمانی **کلان** — شاخص کل، دلار، طلا، تتر و ژئوپلیتیک.

جایگزین طراحی قبلی که چرخه تک‌سهم را می‌سنجید. پرسش این فایل عوض شده:

    نه «چرخه فولاد چند روزه است»
    بلکه «رشد بورس واقعی بود یا فقط دلار؟ دلار چند روز جلوتر حرکت
    می‌کند؟ بازار حول تنش ژئوپلیتیک چه کرد؟»
"""
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as K
from common import (BORDER, F_DATE, F_INT, F_NUM1, F_NUM2, F_PCT, F_PCT2,
                    F_PRICE, FONT, hdr, note, style_data, title_block, widths)

FIRST = 5
N_SPARE = 900

# سری‌های کلان. کلید، برچسب، منبع، و اینکه ماکرو از کجا می‌گیردش.
MACRO_SERIES = [
    ("tedpix",   "شاخص کل بورس",     "tsetmc / شاخص کل"),
    ("tedpix_ew", "شاخص هم‌وزن",      "tsetmc / هم‌وزن"),
    ("usd",      "دلار آزاد",         "tgju / price_dollar_rl"),
    ("nima",     "دلار نیمایی",       "tgju / nima_sell_usd"),
    ("usdt",     "تتر",               "nobitex / usdt-rls"),
    ("ons",      "اونس طلا (دلار)",   "tgju / ons"),
    ("coin",     "سکه تمام",          "tgju / sekee"),
]


def _sheet(wb, name, title, subtitle, ncol):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    title_block(ws, title, subtitle, ncol)
    return ws


# ------------------------------------------------------ ۱) سری‌های کلان
def build_macro_series(wb):
    """جدول خام روزانه. ماکرو با یک کلیک پرش می‌کند."""
    cols = [("تاریخ میلادی", 13, F_DATE), ("تاریخ شمسی", 13, None)]
    for _k, lbl, _s in MACRO_SERIES:
        cols.append((lbl, 15, F_PRICE))
    ws = _sheet(wb, "Macro_Series",
                "سری‌های کلان روزانه",
                "ستون‌ها را دکمه «به‌روزرسانی کلان» پر می‌کند. "
                "هیچ فرمولی اینجا نیست — این شیت فقط داده خام است.",
                len(cols))
    widths(ws, [c[1] for c in cols])
    hdr(ws, 3, [c[0] for c in cols],
        ["-", "-"] + [s for _k, _l, s in MACRO_SERIES])
    for i in range(N_SPARE):
        r = FIRST + i
        for j in range(1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
            c.font = Font(name=FONT, size=8, color=K.C_BLUE_INPUT)
            c.border = BORDER
            if cols[j - 1][2]:
                c.number_format = cols[j - 1][2]
    ws.freeze_panes = "C5"
    return ws


# ------------------------------------------- ۲) شاخص واقعی (به دلار/طلا)
def build_real_index(wb):
    """مهم‌ترین شیت این فایل، و ساده‌ترینش.

    شاخصی که ۴۰٪ رشد کند در حالی که دلار ۵۰٪ رفته، به قدرت خرید **زیان**
    داده. این تعدیل در بیشتر تحلیل‌های داخلی غایب است.
    """
    last = FIRST + N_SPARE - 1
    cols = [
        ("تاریخ", 13, F_DATE, "='Macro_Series'!$A{r}"),
        ("شاخص کل (ریالی)", 15, F_PRICE, "=IFERROR('Macro_Series'!$C{r},\"\")"),
        ("دلار آزاد", 14, F_PRICE, "=IFERROR('Macro_Series'!$E{r},\"\")"),
        ("شاخص به دلار", 15, F_NUM2,
         '=IF(OR(NOT(ISNUMBER($B{r})),NOT(ISNUMBER($C{r})),$C{r}=0),"",$B{r}/$C{r})'),
        ("اونس طلا", 13, F_NUM2, "=IFERROR('Macro_Series'!$H{r},\"\")"),
        ("شاخص به طلا", 15, F_NUM2,
         '=IF(OR(NOT(ISNUMBER($D{r})),NOT(ISNUMBER($E{r})),$E{r}=0),"",$D{r}/$E{r})'),
        ("بازده ریالی ۱ ساله", 14, F_PCT,
         '=IF(OR(NOT(ISNUMBER($B{r})),NOT(ISNUMBER($B{y}))),"",$B{r}/$B{y}-1)'),
        ("بازده دلاری ۱ ساله", 14, F_PCT,
         '=IF(OR(NOT(ISNUMBER($D{r})),NOT(ISNUMBER($D{y}))),"",$D{r}/$D{y}-1)'),
        ("توهم تورمی", 14, F_PCT,
         '=IF(OR(NOT(ISNUMBER($G{r})),NOT(ISNUMBER($H{r}))),"",$G{r}-$H{r})'),
    ]
    ws = _sheet(wb, "Real_Index",
                "شاخص واقعی — بورس به دلار و به طلا",
                "ستون آخر «توهم تورمی» است: چقدر از بازده ریالی، فقط کاهش "
                "ارزش پول بوده و نه رشد واقعی. عدد مثبت بزرگ یعنی بازده "
                "اسمی، واقعی نبوده.", len(cols))
    widths(ws, [c[1] for c in cols])
    hdr(ws, 3, [c[0] for c in cols],
        ["Macro_Series", "ریالی", "ریال", "B÷C", "دلار", "D÷E",
         "۲۴۵ روز", "۲۴۵ روز", "ریالی − دلاری"])
    for i in range(N_SPARE):
        r = FIRST + i
        y = max(FIRST, r - 245)
        for j, (_l, _w, fmt, tmpl) in enumerate(cols, start=1):
            c = ws.cell(row=r, column=j, value=tmpl.format(r=r, y=y))
            c.font = Font(name=FONT, size=8)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER
            if fmt:
                c.number_format = fmt
    ws.conditional_formatting.add("I%d:I%d" % (FIRST, last), ColorScaleRule(
        start_type="num", start_value=-0.5, start_color=K.C_SBUY_BG,
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=0.5, end_color=K.C_SSELL_BG))
    ws.freeze_panes = "B5"
    return ws


# ------------------------------------------------ ۳) رابطه بین سری‌ها
LL_ROWS = [
    ("دلار آزاد", "شاخص کل", "usd", "tedpix"),
    ("اونس طلا", "شاخص کل", "ons", "tedpix"),
    ("تتر", "دلار آزاد", "usdt", "usd"),
    ("دلار آزاد", "سکه تمام", "usd", "coin"),
    ("شاخص کل", "شاخص هم‌وزن", "tedpix", "tedpix_ew"),
    ("دلار نیمایی", "دلار آزاد", "nima", "usd"),
]


def build_lead_lag(wb):
    cols = ["سری پیشرو (فرضی)", "سری پیرو", "تأخیر (روز)", "همبستگی",
            "باند تصحیح‌شده", "باند خام", "مرتبه AR", "تعداد داده",
            "وقفه‌های آزموده", "نتیجه"]
    w = [18, 16, 11, 11, 13, 11, 10, 11, 12, 34]
    ws = _sheet(wb, "Lead_Lag",
                "پیشرو و پیرو — کدام سری زودتر حرکت می‌کند؟",
                "روش باکس-جنکینز با پیش‌سفیدسازی. ⚠️ «باند تصحیح‌شده» بابت "
                "تعداد وقفه‌های آزموده‌شده تعدیل شده؛ ستون «باند خام» فقط "
                "برای نشان دادن اینکه بدون تصحیح چقدر گمراه‌کننده می‌شد.",
                len(cols))
    widths(ws, w)
    hdr(ws, 3, cols, ["ورودی", "ورودی", "پایتون/ماکرو", "−۱ تا ۱",
                      "Šidák", "۱٫۹۶/√n", "AIC", "-", "-", "-"])
    for i, (a, b, _ka, _kb) in enumerate(LL_ROWS):
        r = FIRST + i
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b)
        for j in range(1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.font = Font(name=FONT, size=8, bold=(j <= 2))
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=(j == len(cols)))
            c.border = BORDER
            if j in (4, 5, 6):
                c.number_format = F_NUM2
        ws.row_dimensions[r].height = 24
    r2 = FIRST + len(LL_ROWS) + 1
    for txt in [
        "تأخیر مثبت یعنی سری اول جلوتر است. تأخیر منفی یعنی برعکسِ فرضیه — "
        "و این خودش یافته است، نه خطا.",
        "⚠️ همبستگی، علیت نیست. حتی تأخیر معنادار هم می‌تواند از یک عامل "
        "سوم مشترک بیاید (مثلاً انتظارات تورمی که هر دو را هم‌زمان می‌راند).",
    ]:
        note(ws, "A%d" % r2, txt, size=9, color="404040")
        ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=len(cols))
        ws.row_dimensions[r2].height = 28
        r2 += 1
    return ws


# ------------------------------------------------- ۴) هم‌انباشتگی
CI_ROWS = [
    ("شاخص کل", "دلار آزاد", "tedpix", "usd"),
    ("شاخص کل", "سکه تمام", "tedpix", "coin"),
    ("سکه تمام", "دلار آزاد", "coin", "usd"),
    ("تتر", "دلار آزاد", "usdt", "usd"),
]


def build_cointegration(wb):
    cols = ["سری وابسته", "سری توضیحی", "ضریب بلندمدت β", "آماره ADF",
            "بحرانی ۵٪", "هم‌انباشته؟", "نیمه‌عمر (روز)", "تعداد داده", "تفسیر"]
    w = [15, 15, 14, 12, 12, 13, 14, 11, 40]
    ws = _sheet(wb, "Cointegration",
                "آیا بورس چیزی جز نماینده دلار است؟",
                "آزمون انگل-گرنجر روی لگاریتم سری‌ها. اگر شاخص کل و دلار "
                "هم‌انباشته باشند، «رشد بورس» عمدتاً بازتاب نرخ ارز است و "
                "سیگنال گرفتن از سطح شاخص، دوباره‌شماری همان اطلاعات است.",
                len(cols))
    widths(ws, w)
    hdr(ws, 3, cols, ["ورودی", "ورودی", "OLS روی log", "روی باقی‌مانده",
                      "مک‌کینون", "آماره < بحرانی", "بازگشت به تعادل", "-", "-"])
    for i, (a, b, _ka, _kb) in enumerate(CI_ROWS):
        r = FIRST + i
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b)
        for j in range(1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.font = Font(name=FONT, size=8, bold=(j <= 2))
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=(j == len(cols)))
            c.border = BORDER
            if j in (3, 4, 5, 7):
                c.number_format = F_NUM2
        ws.row_dimensions[r].height = 26
    r2 = FIRST + len(CI_ROWS) + 1
    note(ws, "A%d" % r2,
         "β نزدیک ۱ به‌علاوه هم‌انباشتگی یعنی شاخص و دلار عملاً یک چیزند و "
         "بورس در بلندمدت فقط تورم ارزی را دنبال می‌کند. β کمتر از ۱ یعنی "
         "بورس کمتر از دلار بازده داده — یعنی نگه‌داشتن دلار بهتر بوده.",
         size=9, color="404040")
    ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=len(cols))
    ws.row_dimensions[r2].height = 30
    return ws


# --------------------------------------------- ۵) چرخه سری‌های کلان
def build_macro_cycles(wb):
    cols = ["سری", "طول چرخه غالب (روز)", "نسبت دامنه", "تعداد تکرار",
            "تاریخ آخرین کف", "تاریخ آخرین سقف", "روز از کف",
            "فاز چرخه", "کف بعدی (تخمین)", "وضعیت"]
    w = [18, 18, 12, 12, 14, 14, 11, 11, 15, 30]
    ws = _sheet(wb, "Macro_Cycles",
                "چرخه‌های زمانی متغیرهای کلان",
                "همان آشکارساز طیفی، ولی روی شاخص کل و دلار و طلا — نه روی "
                "تک‌سهم. آشکارساز عمداً سخت‌گیر است: روی نویز خالص صفر بار "
                "چرخه می‌بیند. «پیدا نشد» یک نتیجه است، نه خرابی.",
                len(cols))
    widths(ws, w)
    hdr(ws, 3, cols, ["-", "Goertzel + Šidák", "قله÷میانه", "قطعات",
                      "پیوت", "پیوت", "شمارش", "روز÷طول", "کف+طول", "-"])
    for i, (_k, lbl, _s) in enumerate(MACRO_SERIES):
        r = FIRST + i
        ws.cell(row=r, column=1, value=lbl)
        for j in range(1, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.font = Font(name=FONT, size=8, bold=(j == 1))
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=(j == len(cols)))
            c.border = BORDER
            if j in (3,):
                c.number_format = F_NUM2
            if j in (5, 6, 9):
                c.number_format = F_DATE
        # فاز و کف بعدی از روی ستون‌های نوشته‌شده، فرمول زنده
        ws.cell(row=r, column=7, value=(
            '=IF(OR(NOT(ISNUMBER($E{r})),COUNT(\'Macro_Series\'!$A:$A)=0),"",'
            'MAX(\'Macro_Series\'!$A$5:$A$904)-$E{r})').format(r=r))
        ws.cell(row=r, column=8, value=(
            '=IF(OR(NOT(ISNUMBER($G{r})),NOT(ISNUMBER($B{r})),$B{r}=0),"",'
            'ROUND($G{r}*5/7,0)/$B{r})').format(r=r))
        ws.cell(row=r, column=8).number_format = F_NUM2
        ws.cell(row=r, column=9, value=(
            '=IF(OR(NOT(ISNUMBER($E{r})),NOT(ISNUMBER($B{r}))),"",'
            '$E{r}+ROUND($B{r}*7/5,0))').format(r=r))
        ws.cell(row=r, column=9).number_format = F_DATE
        ws.row_dimensions[r].height = 24
    return ws


# ------------------------------------------------- ۶) رویدادهای ژئوپلیتیک
def build_geo_events(wb):
    cols = ["تاریخ میلادی", "تاریخ شمسی", "رویداد", "دسته",
            "CAR شاخص کل", "CAR دلار", "CAR سکه", "یادداشت"]
    w = [13, 13, 34, 14, 13, 12, 12, 30]
    ws = _sheet(wb, "Geo_Events",
                "رویدادهای ژئوپلیتیک و واکنش بازار",
                "⚠️ **این فهرست را خودتان پر می‌کنید.** من تاریخ رویداد "
                "نمی‌سازم — چیزی که تاریخش را مطمئن نباشم، ننوشتنش بهتر از "
                "نوشتنِ حدسی است. ستون‌های CAR را تحلیل حساب می‌کند.",
                len(cols))
    widths(ws, w)
    hdr(ws, 3, cols, ["ورودی", "خودکار", "ورودی", "ورودی",
                      "±روز پنجره", "±روز پنجره", "±روز پنجره", "ورودی"])
    for i in range(60):
        r = FIRST + i
        for j in (1, 3, 4, 8):
            c = ws.cell(row=r, column=j)
            c.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
            c.font = Font(name=FONT, size=8, color=K.C_BLUE_INPUT)
            c.border = BORDER
        ws.cell(row=r, column=1).number_format = F_DATE
        for j in (2, 5, 6, 7):
            c = ws.cell(row=r, column=j)
            c.font = Font(name=FONT, size=8)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER
            if j > 2:
                c.number_format = F_PCT
    last = FIRST + 59
    for col in ("E", "F", "G"):
        ws.conditional_formatting.add("%s%d:%s%d" % (col, FIRST, col, last),
                                      ColorScaleRule(
            start_type="num", start_value=-0.1, start_color=K.C_SSELL_BG,
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=0.1, end_color=K.C_SBUY_BG))

    r2 = last + 2
    for txt in [
        "CAR = بازده تجمعی غیرعادی در پنجره حول رویداد. **پیش‌فرض [−۱، +۳] روز** — با اندازه‌گیری انتخاب شد نه با عرف: پنجره پهن‌تر ([−۵، +۱۰]) در آزمون، شوک واقعی ۳٫۵٪ را زیر نویز گم کرد (p=۰٫۴۹)، در حالی که پنجره باریک همان شوک را با p=۰٫۰۰۴ گرفت.",
        "معناداری با **آزمون جایگشت** سنجیده می‌شود: همان تعداد تاریخ ولی "
        "تصادفی. اگر CAR واقعی از توزیع تاریخ‌های تصادفی جدا نباشد، رویدادها "
        "اطلاعاتی نداشته‌اند. آزمون پارامتری اینجا نامناسب است چون بازده‌ها "
        "دم‌سنگین‌اند و تعداد رویداد کم.",
        "⚠️ با کمتر از حدود ۱۰ رویداد، هیچ نتیجه آماری‌ای معتبر نیست — "
        "جدول را می‌شود توصیفی خواند، نه استنباطی.",
    ]:
        note(ws, "A%d" % r2, txt, size=9, color="404040")
        ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=len(cols))
        ws.row_dimensions[r2].height = 30
        r2 += 1
    ws.freeze_panes = "C5"
    return ws
