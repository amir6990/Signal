# -*- coding: utf-8 -*-
"""
Gold_Analysis.xlsx — تحلیل طلا و سکه.

شیت‌ها: Gold_Dashboard · Gold_Input · Gold_History · Coin_Bubble · Settings · Documentation

منطق مرکزی این فایل: قیمت سکه در ایران دو جزء دارد که باید از هم جدا شوند —
ارزش ذاتی طلای داخل آن، و حباب. بازده ریالی طلا اغلب نه از طلا، بلکه از
تضعیف ریال یا صرفاً باد کردن حباب می‌آید. فایل هر سه محرک را تفکیک می‌کند.
"""
import datetime
import random

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

import build_workbook as B
import common as K
from common import (BORDER, F_DATE, F_NUM1, F_NUM2, F_PCT, F_PCT2, F_PRICE, FONT,
                    hdr, note, title_block, widths)
from .assets_common import build_history_sheet, trend_block

SHEET_ORDER = ["Gold_Dashboard", "Coin_Bubble", "Gold_Input", "Gold_History",
               "Settings", "Documentation"]

OUNCE_GRAMS = 31.1034768
COINS = [
    ("سکه تمام بهار آزادی", 8.133, 0.900),
    ("نیم سکه", 4.0665, 0.900),
    ("ربع سکه", 2.03325, 0.900),
    ("سکه گرمی", 1.0166, 0.900),
    ("طلای ۱۸ عیار (هر گرم)", 1.0, 0.750),
]

GOLD_DOC = [
    ("H1", "راهنمای فایل تحلیل طلا و سکه"),
    ("W", "⚠️ داده‌های داخل فایل نمونه (DEMO) و ساختگی‌اند. قیمت‌های واقعی را در "
          "شیت Gold_Input و Gold_History وارد کنید."),
    ("H2", "۱) منطق تفکیک قیمت"),
    ("P", "قیمت سکه = ارزش ذاتی طلای داخل آن + حباب. ارزش ذاتی از اونس جهانی و "
          "نرخ ارز به‌دست می‌آید و هیچ فرضی ندارد — یک محاسبه قطعی است."),
    ("T", "ارزش ذاتی|طلای خالص (گرم) × (اونس ÷ ۳۱٫۱۰۳۴۷۶۸) × نرخ ارز"),
    ("T", "حباب|قیمت بازار ÷ ارزش ذاتی − ۱"),
    ("P", "سکه تمام بهار آزادی: ۸٫۱۳۳ گرم با عیار ۹۰۰ ← ۷٫۳۱۹۷ گرم طلای خالص."),
    ("H2", "۲) تفکیک بازده به سه محرک"),
    ("P", "بازده ریالی طلا سه منشأ دارد و شیت Coin_Bubble هر سه را جدا می‌کند: "
          "تغییر اونس جهانی، تغییر نرخ ارز، و تغییر حباب. این تفکیک برای بازار "
          "ایران کلیدی است — رشد ریالی طلا اغلب اصلاً درباره طلا نیست."),
    ("T", "کنترل صحت|(۱+بازده اونس) × (۱+بازده ارز) × (۱+تغییر حباب) − ۱ = بازده کل"),
    ("P", "اگر این کنترل با بازده کل نخواند، یعنی یکی از ورودی‌ها ناسازگار است."),
    ("H2", "۳) حباب چه می‌گوید و چه نمی‌گوید"),
    ("P", "حباب بالا یعنی بازار داخلی نسبت به ارزش ذاتی گران است — ولی حباب "
          "می‌تواند ماه‌ها بالا بماند یا بالاتر برود. حباب یک سنجه ارزش‌گذاری "
          "است، نه سیگنال زمان‌بندی. تاریخچه حباب خودتان را در Gold_History "
          "ثبت کنید تا بدانید عدد امروز نسبت به گذشته کجاست."),
    ("H2", "۴) آنچه این فایل نمی‌کند"),
    ("L", "• چرخه ادعایی «۸ ساله طلا» عمداً کدگذاری نشده — در منابع این پروژه "
          "مستند نیست. اگر چنین چرخه‌ای هست، باید در خروجی موتور طیفی روی سری "
          "واقعی ظاهر شود: python -m timeframe symbol --csv gold.csv --name طلا"),
    ("L", "• اجرت ساخت، مالیات و سود فروشنده در ارزش ذاتی لحاظ نشده؛ عدد محاسبه‌شده "
          "ارزش خام فلز است."),
    ("L", "• برای طلای آبشده و مسکوکات دیگر، وزن و عیار را در شیت Coin_Bubble "
          "تغییر دهید."),
]


def _sample_history(n=500):
    rng = random.Random(21)
    d = datetime.date(2026, 9, 3)
    days = []
    while len(days) < n:
        if d.weekday() not in (3, 4):
            days.append(d)
        d -= datetime.timedelta(days=1)
    days.reverse()
    p, out = 620_000_000.0, []
    for dd in days:
        p *= 1 + rng.gauss(0.0009, 0.013)
        hi, lo = p * (1 + abs(rng.gauss(0, 0.004))), p * (1 - abs(rng.gauss(0, 0.004)))
        out.append((dd, round(p), round(hi), round(lo), 0))
    return out


def _input_sheet(wb):
    ws = wb.create_sheet("Gold_Input")
    ws.sheet_view.rightToLeft = True
    title_block(ws, "ورودی‌های لحظه‌ای طلا",
                "سلول‌های زرد را پر کنید. همه محاسبات فایل از همین چند عدد می‌آیند.", 4)
    widths(ws, [40, 22, 46, 3])
    rows = [
        ("S", "بازار جهانی", None, None, None),
        ("P", "اونس طلا (دلار)", 3400.0, "GOLD_OZ", "منبع: بازار جهانی"),
        ("P", "اونس نقره (دلار)", 40.0, "SILVER_OZ", "اختیاری — برای نسبت طلا به نقره"),
        ("B", None, None, None, None),
        ("S", "نرخ ارز", None, None, None),
        ("P", "دلار آزاد (ریال)", 950_000.0, "USD_FREE", "نرخ بازار آزاد"),
        ("P", "دلار نیمایی / مرجع (ریال)", 720_000.0, "USD_NIMA", "برای محاسبه شکاف"),
        ("B", None, None, None, None),
        ("S", "قیمت بازار داخلی", None, None, None),
        ("P", "سکه تمام بهار آزادی (ریال)", 900_000_000.0, "COIN_FULL", "قیمت بازار"),
        ("P", "نیم سکه (ریال)", 470_000_000.0, "COIN_HALF", ""),
        ("P", "ربع سکه (ریال)", 285_000_000.0, "COIN_QTR", ""),
        ("P", "طلای ۱۸ عیار هر گرم (ریال)", 82_000_000.0, "GRAM_18K", "شامل اجرت و مالیات"),
        ("B", None, None, None, None),
        ("S", "مبنای مقایسه (برای تفکیک بازده)", None, None, None),
        ("P", "اونس در تاریخ مبنا (دلار)", 2600.0, "GOLD_OZ_0", "مثلاً ابتدای سال"),
        ("P", "دلار در تاریخ مبنا (ریال)", 600_000.0, "USD_FREE_0", ""),
        ("P", "سکه در تاریخ مبنا (ریال)", 500_000_000.0, "COIN_FULL_0", ""),
    ]
    r = 4
    for kind, label, val, name, desc in rows:
        if kind == "B":
            r += 1
            continue
        if kind == "S":
            ws.cell(row=r, column=1, value=label)
            for c in (1, 2, 3):
                cell = ws.cell(row=r, column=c)
                cell.fill = PatternFill("solid", fgColor=K.C_SECTION_BG)
                cell.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
                cell.alignment = Alignment(horizontal="right", vertical="center")
            ws.row_dimensions[r].height = 20
            r += 1
            continue
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=9)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
        c = ws.cell(row=r, column=2, value=val)
        c.font = Font(name=FONT, size=10, bold=True, color=K.C_BLUE_INPUT)
        c.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        c.number_format = F_NUM2 if val and val < 1e4 else '#,##0'
        ws.cell(row=r, column=3, value=desc).font = Font(
            name=FONT, size=8, italic=True, color=K.C_NOTE)
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="right", vertical="center")
        B.dn(name, "Gold_Input", "$B$%d" % r)
        r += 1
    note(ws, "A%d" % (r + 1),
         "⚠️ اعداد بالا نمونه‌اند. قیمت واقعی روز را جایگزین کنید — همه خروجی‌های "
         "فایل مستقیماً از همین سلول‌ها می‌آیند.", color="C00000", bold=True, size=9)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)
    ws.freeze_panes = "A4"
    return ws


def _bubble_sheet(wb):
    ws = wb.create_sheet("Coin_Bubble")
    ws.sheet_view.rightToLeft = True
    cols = ["مسکوک", "وزن (گرم)", "عیار", "طلای خالص (گرم)",
            "ارزش ذاتی (ریال)", "قیمت بازار (ریال)", "حباب (ریال)", "حباب ٪"]
    title_block(ws, "ارزش ذاتی و حباب مسکوکات",
                "ارزش ذاتی یک محاسبه قطعی است: محتوای طلای خالص × قیمت جهانی × نرخ ارز. "
                "حباب همان اختلاف قیمت بازار با این عدد است.", len(cols))
    widths(ws, [26, 12, 10, 15, 20, 20, 20, 12])
    hdr(ws, 3, cols, ["-", "ورودی", "ورودی", "وزن × عیار",
                      "GOLD_OZ · USD_FREE", "ورودی", "بازار − ذاتی", "÷ ذاتی − ۱"])
    market_names = ["COIN_FULL", "COIN_HALF", "COIN_QTR", "", "GRAM_18K"]
    r = 5
    for i, (name, w, purity) in enumerate(COINS):
        ws.cell(row=r, column=1, value=name).font = Font(name=FONT, size=9, bold=True)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
        for col, val in ((2, w), (3, purity)):
            c = ws.cell(row=r, column=col, value=val)
            c.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
            c.font = Font(name=FONT, size=9, color=K.C_BLUE_INPUT)
            c.number_format = F_NUM2 if col == 2 else '0.000'
        ws.cell(row=r, column=4, value='=$B{r}*$C{r}'.format(r=r)).number_format = '0.0000'
        ws.cell(row=r, column=5,
                value='=$D{r}*(GOLD_OZ/{g})*USD_FREE'.format(r=r, g=OUNCE_GRAMS)
                ).number_format = '#,##0'
        mk = market_names[i]
        if mk:
            ws.cell(row=r, column=6, value="=%s" % mk).number_format = '#,##0'
        else:
            c = ws.cell(row=r, column=6)
            c.fill = PatternFill("solid", fgColor=K.C_INPUT_BG)
            c.font = Font(name=FONT, size=9, color=K.C_BLUE_INPUT)
            c.number_format = '#,##0'
        ws.cell(row=r, column=7,
                value='=IFERROR($F{r}-$E{r},"")'.format(r=r)).number_format = '#,##0'
        ws.cell(row=r, column=8,
                value='=IFERROR($F{r}/$E{r}-1,"")'.format(r=r)).number_format = F_PCT
        for cc in range(1, len(cols) + 1):
            ws.cell(row=r, column=cc).border = BORDER
            ws.cell(row=r, column=cc).alignment = Alignment(
                horizontal="right" if cc == 1 else "center", vertical="center")
            if cc >= 4:
                ws.cell(row=r, column=cc).font = Font(name=FONT, size=9)
        r += 1
    last = r - 1
    ws.conditional_formatting.add("H5:H%d" % last, ColorScaleRule(
        start_type="num", start_value=-0.05, start_color=K.C_SBUY_BG,
        mid_type="num", mid_value=0.15, mid_color="FFFFFF",
        end_type="num", end_value=0.45, end_color=K.C_SSELL_BG))
    B.dn("BUBBLE_FULL", "Coin_Bubble", "$H$5")

    r += 2
    ws.cell(row=r, column=1, value="تفکیک بازده ریالی سکه به سه محرک")
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = PatternFill("solid", fgColor=K.C_SECTION_BG)
        cell.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        cell.alignment = Alignment(horizontal="right", vertical="center")
    r += 1
    decomp = [
        ("بازده کل ریالی سکه", '=IFERROR(COIN_FULL/COIN_FULL_0-1,"")', F_PCT),
        ("سهم طلای جهانی", '=IFERROR(GOLD_OZ/GOLD_OZ_0-1,"")', F_PCT),
        ("سهم نرخ ارز", '=IFERROR(USD_FREE/USD_FREE_0-1,"")', F_PCT),
        ("حباب مبنا", '=IFERROR(COIN_FULL_0/({d}*(GOLD_OZ_0/{g})*USD_FREE_0)-1,"")'
         .format(d=COINS[0][1] * COINS[0][2], g=OUNCE_GRAMS), F_PCT),
        ("حباب کنونی", '=IFERROR($H$5,"")', F_PCT),
        ("سهم تغییر حباب", '=IFERROR((1+$B{c})/(1+$B{b})-1,"")', F_PCT),
        ("کنترل صحت (باید با بازده کل یکی باشد)",
         '=IFERROR((1+$B{g})*(1+$B{f})*(1+$B{p})-1,"")', F_PCT),
    ]
    base = r
    for i, (lbl, f, fmt) in enumerate(decomp):
        ws.cell(row=r, column=1, value=lbl).font = Font(
            name=FONT, size=9, bold=(i in (0, 6)))
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
        if "{" in f and ("{c}" in f or "{b}" in f or "{g}" in f):
            f = f.format(c=base + 4, b=base + 3, g=base + 1, f=base + 2, p=base + 5)
        c = ws.cell(row=r, column=2, value=f)
        c.font = Font(name=FONT, size=10, bold=True, color=K.C_GREEN_LINK)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        r += 1
    ws.cell(row=r, column=1, value="اختلاف کنترل با بازده کل")
    ws.cell(row=r, column=1).font = Font(name=FONT, size=9, italic=True)
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="center")
    chk = ws.cell(row=r, column=2, value='=IFERROR($B{t}-$B{b},"")'.format(
        t=base + 6, b=base))
    chk.number_format = '0.0000%'
    chk.font = Font(name=FONT, size=9, bold=True)
    chk.border = BORDER
    ws.conditional_formatting.add("B%d" % r, CellIsRule(
        operator="between", formula=["-0.0001", "0.0001"],
        fill=PatternFill("solid", fgColor=K.C_BUY_BG), font=Font(color=K.C_BUY_FG, bold=True)))
    ws.conditional_formatting.add("B%d" % r, CellIsRule(
        operator="notBetween", formula=["-0.0001", "0.0001"],
        fill=PatternFill("solid", fgColor=K.C_SSELL_BG), font=Font(color="FFFFFF", bold=True)))
    note(ws, "A%d" % (r + 2),
         "چرا این تفکیک مهم است: رشد ریالی طلا در ایران اغلب اصلاً درباره طلا نیست. "
         "اگر «سهم نرخ ارز» بزرگ‌ترین جزء باشد، شما روی ریال شرط بسته‌اید، نه روی طلا.")
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=len(cols))
    ws.freeze_panes = "A5"
    return ws


def _dashboard(wb, hist_last, trend_row):
    ws = wb.create_sheet("Gold_Dashboard", 0)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    widths(ws, [34, 22, 4, 34, 22])
    title_block(ws, "داشبورد طلا و سکه",
                "⚠️ داده‌های نمونه‌اند. ورودی‌ها را در شیت Gold_Input به‌روز کنید.", 5)
    ws["A2"].font = Font(name=FONT, size=9, bold=True, italic=True, color="C00000")

    def section(row, text, c2=5):
        ws.cell(row=row, column=1, value=text)
        for c in range(1, c2 + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = PatternFill("solid", fgColor=K.C_SECTION_BG)
            cell.font = Font(name=FONT, bold=True, size=11, color="FFFFFF")
            cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row].height = 22

    def kpi(row, col, label, formula, fmt=None, big=False):
        ws.cell(row=row, column=col, value=label).font = Font(name=FONT, size=9)
        ws.cell(row=row, column=col).alignment = Alignment(
            horizontal="right", vertical="center")
        c = ws.cell(row=row, column=col + 1, value=formula)
        c.font = Font(name=FONT, size=12 if big else 10, bold=True, color=K.C_GREEN_LINK)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        if fmt:
            c.number_format = fmt
        return c

    section(4, "۱) قیمت‌ها و ارزش‌گذاری")
    kpi(5, 1, "اونس جهانی (دلار)", "=GOLD_OZ", '#,##0.00')
    kpi(6, 1, "دلار آزاد (ریال)", "=USD_FREE", '#,##0')
    kpi(7, 1, "سکه تمام — بازار", "=COIN_FULL", '#,##0')
    kpi(8, 1, "سکه تمام — ارزش ذاتی", "=Coin_Bubble!$E$5", '#,##0')
    kpi(9, 1, "حباب سکه تمام", "=BUBBLE_FULL", F_PCT, big=True)
    ws.conditional_formatting.add("B9", ColorScaleRule(
        start_type="num", start_value=-0.05, start_color=K.C_SBUY_BG,
        mid_type="num", mid_value=0.15, mid_color="FFFFFF",
        end_type="num", end_value=0.45, end_color=K.C_SSELL_BG))

    kpi(5, 4, "شکاف دلار آزاد و نیمایی", '=IFERROR(USD_FREE/USD_NIMA-1,"")', F_PCT)
    kpi(6, 4, "نسبت طلا به نقره", '=IFERROR(GOLD_OZ/SILVER_OZ,"")', F_NUM1)
    kpi(7, 4, "ارزش خام گرم ۱۸ عیار", "=Coin_Bubble!$E$9", '#,##0')
    kpi(8, 4, "حباب گرم ۱۸ عیار", "=Coin_Bubble!$H$9", F_PCT)
    kpi(9, 4, "حباب ربع سکه", "=Coin_Bubble!$H$7", F_PCT)

    section(11, "۲) روند قیمت سکه (از شیت Gold_History)")
    lbl_row, score_row = trend_block(ws, "Gold_History", hist_last, 12, "سکه")

    r = score_row + 3
    section(r, "۳) تفکیک بازده — این رشد از کجا آمده؟")
    r += 1
    for lbl, ref in (("بازده کل ریالی", "$B$%d" % trend_row),
                     ("سهم طلای جهانی", "$B$%d" % (trend_row + 1)),
                     ("سهم نرخ ارز", "$B$%d" % (trend_row + 2)),
                     ("سهم تغییر حباب", "$B$%d" % (trend_row + 5))):
        kpi(r, 1, lbl, "=Coin_Bubble!%s" % ref, F_PCT)
        r += 1
    note(ws, "A%d" % (r + 1),
         "اگر «سهم نرخ ارز» بزرگ‌ترین جزء باشد، این یک موقعیت ارزی است که لباس طلا "
         "پوشیده. حباب هم می‌تواند ماه‌ها بالا بماند — سنجه ارزش‌گذاری است، نه زمان‌بندی.")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
    ws.freeze_panes = "A4"
    return ws


def build(out_path, **_kw):
    B.DEFINED.clear()
    B.COLREF.clear()
    wb = Workbook()
    wb.remove(wb.active)

    B.build_settings(wb, sections=("۱)", "۲)", "۸)"))
    _input_sheet(wb)
    bubble = _bubble_sheet(wb)
    rows = _sample_history()
    _, hist_last = build_history_sheet(
        wb, "Gold_History", "تاریخچه روزانه سکه تمام بهار آزادی",
        "ستون‌های A تا F ورودی؛ بقیه فرمول. ردیف‌های زرد خالی ظرفیت آماده‌اند.", rows)
    # ردیف شروع بلوک تفکیک در Coin_Bubble
    trend_row = 5 + len(COINS) + 2 + 1
    _dashboard(wb, hist_last, trend_row)
    B.build_documentation(wb, scope="minimal", extra=GOLD_DOC)

    for name, ref in B.DEFINED.items():
        wb.defined_names.add(DefinedName(name, attr_text=ref))
    order = {n: i for i, n in enumerate(SHEET_ORDER)}
    wb._sheets.sort(key=lambda s: order.get(s.title, 99))
    wb.active = 0
    wb.save(out_path)
    print("✓ %s — %d شیت" % (out_path, len(wb.sheetnames)))
    return out_path
