# -*- coding: utf-8 -*-
"""ظرف سری زمانی و بارگذاری آن از اکسل یا CSV."""
import csv
import datetime
from dataclasses import dataclass, field
from typing import List, Optional, Sequence


@dataclass
class Bar:
    date: datetime.date
    close: float
    high: Optional[float] = None
    low: Optional[float] = None
    open: Optional[float] = None
    volume: Optional[float] = None


@dataclass
class TimeSeries:
    """سری زمانی مرتب‌شده صعودی بر اساس تاریخ."""
    name: str
    bars: List[Bar] = field(default_factory=list)

    def __post_init__(self):
        self.bars.sort(key=lambda b: b.date)

    def __len__(self):
        return len(self.bars)

    @property
    def dates(self) -> List[datetime.date]:
        return [b.date for b in self.bars]

    @property
    def closes(self) -> List[float]:
        return [b.close for b in self.bars]

    @property
    def highs(self) -> List[float]:
        return [b.high if b.high is not None else b.close for b in self.bars]

    @property
    def lows(self) -> List[float]:
        return [b.low if b.low is not None else b.close for b in self.bars]

    @property
    def last_date(self) -> datetime.date:
        return self.bars[-1].date

    @property
    def last_close(self) -> float:
        return self.bars[-1].close

    def index_of(self, d: datetime.date) -> Optional[int]:
        """نزدیک‌ترین اندیس به تاریخ داده‌شده (کندل در تاریخ یا قبل از آن)."""
        lo, hi, out = 0, len(self.bars) - 1, None
        while lo <= hi:
            mid = (lo + hi) // 2
            if self.bars[mid].date <= d:
                out, lo = mid, mid + 1
            else:
                hi = mid - 1
        return out

    def slice_last(self, n: int) -> "TimeSeries":
        return TimeSeries(self.name, list(self.bars[-n:]))

    def bars_between(self, d1: datetime.date, d2: datetime.date) -> int:
        """تعداد کندل معاملاتی بین دو تاریخ (نه روز تقویمی)."""
        i1, i2 = self.index_of(d1), self.index_of(d2)
        if i1 is None or i2 is None:
            return 0
        return abs(i2 - i1)

    def bars_per_calendar_day(self) -> float:
        """نسبت تجربی کندل به روز تقویمی — برای تبدیل چرخه‌های تقویمی به کندل."""
        if len(self.bars) < 2:
            return 1.0
        span = (self.bars[-1].date - self.bars[0].date).days
        return (len(self.bars) - 1) / span if span > 0 else 1.0


# ------------------------------------------------------------------ loaders
def from_csv(path: str, name: str, date_col="date", close_col="close",
             high_col="high", low_col="low", volume_col="volume",
             date_format="%Y-%m-%d") -> TimeSeries:
    bars = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            try:
                d = datetime.datetime.strptime(row[date_col].strip(), date_format).date()
                c = float(row[close_col])
            except (KeyError, ValueError, AttributeError):
                continue

            def _f(col):
                try:
                    return float(row[col])
                except (KeyError, ValueError, TypeError):
                    return None

            bars.append(Bar(d, c, _f(high_col), _f(low_col), None, _f(volume_col)))
    return TimeSeries(name, bars)


def from_workbook(path: str, symbol: str, sheet="Daily_History",
                  col_symbol=1, col_date=3, col_high=7, col_low=8,
                  col_close=9, col_volume=12, first_row=5) -> TimeSeries:
    """خواندن تاریخچه یک نماد از شیت Daily_History فایل سیگنال.

    شماره ستون‌ها با ساختار build_workbook.py هماهنگ است (A=1).
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    bars = []
    for row in ws.iter_rows(min_row=first_row, values_only=True):
        if not row or row[col_symbol - 1] != symbol:
            continue
        d = row[col_date - 1]
        c = row[col_close - 1]
        if d is None or c is None:
            continue
        if isinstance(d, datetime.datetime):
            d = d.date()
        bars.append(Bar(d, float(c),
                        _num(row, col_high), _num(row, col_low),
                        None, _num(row, col_volume)))
    wb.close()
    return TimeSeries(symbol, bars)


def workbook_symbols(path: str, sheet="Daily_History", col_symbol=1, first_row=5) -> List[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    seen, out = set(), []
    for row in ws.iter_rows(min_row=first_row, min_col=col_symbol,
                            max_col=col_symbol, values_only=True):
        s = row[0]
        if s and s not in seen:
            seen.add(s)
            out.append(s)
    wb.close()
    return out


def _num(row: Sequence, col: int) -> Optional[float]:
    try:
        v = row[col - 1]
        return float(v) if v is not None else None
    except (IndexError, TypeError, ValueError):
        return None


def synthetic(name: str, n: int, periods_amps, noise=0.0, trend=0.0,
              base=100.0, seed=7, start=datetime.date(2020, 1, 1)) -> TimeSeries:
    """سری مصنوعی با چرخه‌های معلوم — فقط برای تست موتور."""
    import math
    import random
    rng = random.Random(seed)
    bars = []
    for i in range(n):
        v = base + trend * i
        for p, a, ph in periods_amps:
            v += a * math.sin(2 * math.pi * (i / p) + ph)
        if noise:
            v += rng.gauss(0, noise)
        bars.append(Bar(start + datetime.timedelta(days=i), v, v * 1.01, v * 0.99))
    return TimeSeries(name, bars)
