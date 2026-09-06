# -*- coding: utf-8 -*-
"""
پر کردن پانل آماری کلان در Time_Analysis.

    python scripts/macro_panel.py --workbook Time_Analysis.xlsx

سه شیت را حساب می‌کند: Lead_Lag، Cointegration، و ستون‌های CAR در Geo_Events.

چرا این‌ها با دکمه اکسل حساب نمی‌شوند: روابط بلندمدت (هم‌انباشتگی،
پیشرو/پیرو) **کند** تغییر می‌کنند. اجرای ماهانه‌شان کافی است و درست‌تر هم
هست — نگاه‌کردن روزانه به یک رابطه بلندمدت، فقط نویز را دنبال کردن است.
داده روزانه و چرخه‌ها را دکمه به‌روز نگه می‌دارد.
"""
import argparse
import datetime
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "src"))

from openpyxl import load_workbook                                  # noqa: E402
from timeframe.macro import series_analysis as SA                   # noqa: E402

# ستون‌های Macro_Series
COLS = {"tedpix": 3, "tedpix_ew": 4, "usd": 5, "nima": 6,
        "usdt": 7, "ons": 8, "coin": 9}
LL_ROWS = [("usd", "tedpix"), ("ons", "tedpix"), ("usdt", "usd"),
           ("usd", "coin"), ("tedpix", "tedpix_ew"), ("nima", "usd")]
CI_ROWS = [("tedpix", "usd"), ("tedpix", "coin"), ("coin", "usd"), ("usdt", "usd")]


def read_series(ws, first=5, last=904):
    dates, cols = [], {k: [] for k in COLS}
    for r in range(first, last + 1):
        d = ws.cell(row=r, column=1).value
        if not isinstance(d, (datetime.date, datetime.datetime)):
            continue
        if isinstance(d, datetime.datetime):
            d = d.date()
        dates.append(d)
        for k, c in COLS.items():
            v = ws.cell(row=r, column=c).value
            cols[k].append(float(v) if isinstance(v, (int, float)) and v > 0 else None)
    return dates, cols


def _pair(dates, cols, a, b):
    d, x, y = [], [], []
    for i, dt in enumerate(dates):
        va, vb = cols[a][i], cols[b][i]
        if va is not None and vb is not None:
            d.append(dt)
            x.append(va)
            y.append(vb)
    return d, x, y


def main(argv=None):
    ap = argparse.ArgumentParser(description="پانل آماری کلان")
    ap.add_argument("--workbook", default="Time_Analysis.xlsx")
    ap.add_argument("--max-lag", type=int, default=30, dest="max_lag")
    ap.add_argument("--pre", type=int, default=1,
                    help="روز قبل از رویداد (پیش‌فرض ۱ — پنجره پهن شوک را گم می‌کند)")
    ap.add_argument("--post", type=int, default=3)
    args = ap.parse_args(argv)

    if not os.path.exists(args.workbook):
        print("فایل یافت نشد: %s" % args.workbook)
        return 1
    wb = load_workbook(args.workbook)
    if "Macro_Series" not in wb.sheetnames:
        print("شیت Macro_Series نیست. اول دکمه «به‌روزرسانی کلان» را بزنید.")
        return 1
    dates, cols = read_series(wb["Macro_Series"])
    n_ok = sum(1 for v in cols["tedpix"] if v)
    print("محور تاریخ: %d روز | شاخص کل: %d مقدار" % (len(dates), n_ok))
    if n_ok < 80:
        print("داده کافی نیست. اول دکمه «به‌روزرسانی کلان» را در اکسل بزنید.")
        return 1

    # ---------------- پیشرو / پیرو ----------------
    ws = wb["Lead_Lag"]
    print("\nپیشرو و پیرو:")
    for i, (a, b) in enumerate(LL_ROWS):
        r = 5 + i
        _d, x, y = _pair(dates, cols, a, b)
        res = SA.prewhitened_ccf(SA.log_returns(x), SA.log_returns(y),
                                 max_lag=args.max_lag) if len(x) > 80 else None
        if res is None:
            ws.cell(row=r, column=10, value="داده کافی نیست")
            print("  %-12s → %-12s داده کافی نیست" % (a, b))
            continue
        ws.cell(row=r, column=3, value=res.lag)
        ws.cell(row=r, column=4, value=round(res.corr, 4))
        ws.cell(row=r, column=5, value=round(res.band, 4))
        ws.cell(row=r, column=6, value=round(res.band_naive, 4))
        ws.cell(row=r, column=7, value=res.ar_order)
        ws.cell(row=r, column=8, value=res.n)
        ws.cell(row=r, column=9, value=res.n_lags_tested)
        if res.significant:
            if res.lag > 0:
                txt = "«%s» حدود %d روز جلوتر حرکت می‌کند" % (a, res.lag)
            elif res.lag < 0:
                txt = "برعکس فرضیه: «%s» %d روز جلوتر است" % (b, -res.lag)
            else:
                txt = "هم‌زمان حرکت می‌کنند، بدون تقدم"
        else:
            txt = "تأخیر معناداری نیست (حتی بیشینه قله از باند رد نشد)"
        ws.cell(row=r, column=10, value=txt)
        print("  %-12s → %-12s lag=%+3d r=%+.2f %s"
              % (a, b, res.lag, res.corr, "✔" if res.significant else "—"))

    # ---------------- هم‌انباشتگی ----------------
    ws = wb["Cointegration"]
    print("\nهم‌انباشتگی:")
    for i, (a, b) in enumerate(CI_ROWS):
        r = 5 + i
        _d, x, y = _pair(dates, cols, a, b)
        res = SA.engle_granger(x, y) if len(x) > 80 else None
        if res is None:
            ws.cell(row=r, column=9, value="داده کافی نیست")
            print("  %-10s ~ %-10s داده کافی نیست" % (a, b))
            continue
        ws.cell(row=r, column=3, value=round(res.beta, 4))
        ws.cell(row=r, column=4, value=round(res.adf_stat, 3))
        ws.cell(row=r, column=5, value=round(res.crit_5, 3))
        ws.cell(row=r, column=6, value="بله" if res.cointegrated else "خیر")
        if res.half_life:
            ws.cell(row=r, column=7, value=round(res.half_life, 1))
        ws.cell(row=r, column=8, value=res.n)
        if res.cointegrated:
            if abs(res.beta - 1.0) < 0.15:
                txt = ("رابطه یک‌به‌یک بلندمدت. «%s» در بلندمدت چیزی جز "
                       "بازتاب «%s» نیست." % (a, b))
            elif res.beta < 1:
                txt = ("هم‌انباشته ولی با ضریب %.2f — «%s» کمتر از «%s» "
                       "بازده داده." % (res.beta, a, b))
            else:
                txt = "هم‌انباشته با ضریب %.2f — بازده بیشتر از %s." % (res.beta, b)
        else:
            txt = "رابطه تعادلی بلندمدت تأیید نشد — دو سری مستقل رفتار می‌کنند."
        ws.cell(row=r, column=9, value=txt)
        print("  %-10s ~ %-10s β=%.3f ADF=%.2f %s"
              % (a, b, res.beta, res.adf_stat,
                 "هم‌انباشته" if res.cointegrated else "نه"))

    # ---------------- مطالعه رویداد ----------------
    if "Geo_Events" in wb.sheetnames:
        we = wb["Geo_Events"]
        events = []
        for r in range(5, 65):
            d = we.cell(row=r, column=1).value
            lbl = we.cell(row=r, column=3).value
            if isinstance(d, (datetime.date, datetime.datetime)):
                events.append((r, lbl or "?",
                               d.date() if isinstance(d, datetime.datetime) else d))
        print("\nمطالعه رویداد: %d رویداد ثبت‌شده" % len(events))
        if events:
            from timeframe.jalali import jalali_str
            for r, lbl, d in events:
                we.cell(row=r, column=2, value=jalali_str(d))
            for key, col in (("tedpix", 5), ("usd", 6), ("coin", 7)):
                d2, x2, _ = _pair(dates, cols, key, key)
                if len(x2) < 80:
                    continue
                es = SA.event_study(d2, x2, [(l, dd) for _r, l, dd in events],
                                    pre=args.pre, post=args.post)
                by = {e.date: e for e in es.events}
                for r, _l, dd in events:
                    e = by.get(dd)
                    if e and e.ok:
                        we.cell(row=r, column=col, value=round(e.car, 4))
                print("  %-8s میانگین CAR=%+.2f%%  p=%.3f  (%d رویداد قابل استفاده)"
                      % (key, es.mean_car * 100, es.p_value, es.n_used))
                if es.n_used < 10:
                    print("      ⚠ کمتر از ۱۰ رویداد — این عدد توصیفی است، نه استنباطی.")
        else:
            print("  هیچ رویدادی در شیت Geo_Events ثبت نشده. "
                  "تاریخ و عنوان را خودتان وارد کنید.")

    wb.save(args.workbook)
    print("\n✓ %s به‌روز شد." % args.workbook)
    print("  فایل را در اکسل باز کنید تا فرمول‌ها دوباره حساب شوند.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
