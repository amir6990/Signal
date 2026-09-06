# -*- coding: utf-8 -*-
"""
پل داده‌ای بین فایل‌های اکسل مجموعه.

    python scripts/link_workbooks.py --stocks Stocks_Signals.xlsx --options Options_Signals.xlsx

چرا این اسکریپت به‌جای ارجاع بین‌فایلی اکسل:
فرمول `='[1]Sheet'!A1` وقتی فایل مبدأ باز نباشد فقط مقدار کش‌شده را نشان
می‌دهد، مسیر فایل را مطلق ذخیره می‌کند (با جابه‌جایی پوشه می‌شکند)، و مهم‌تر:
با بازنویسی فایل توسط openpyxl کاملاً از بین می‌رود. کپی صریح داده پایدارتر
و قابل‌ردگیری‌تر است — و تاریخ به‌روزرسانی هم ثبت می‌شود.
"""
import argparse
import datetime
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))

from openpyxl import load_workbook   # noqa: E402

import build_workbook as BW            # noqa: E402

# بازه‌های واقعی داده. خواندن تا max_row اشتباه است: متن پانویس زیر جدول هم
# به‌عنوان نماد خوانده می‌شود — همان باگی که قبلاً در Watchlist رخ داد.
SIGNALS_FIRST, SIGNALS_ROWS = 4, BW.N_SYM_ROWS
UNDER_FIRST, UNDER_ROWS = 5, BW.UNDERLYING_ROWS


def _read_stocks(path):
    """نماد → (آخرین قیمت، قیمت پایانی، امتیاز کل، سیگنال، قدرت)."""
    wb = load_workbook(path, data_only=True)
    if "Signals" not in wb.sheetnames:
        raise SystemExit("شیت Signals در %s نیست." % path)
    ws = wb["Signals"]
    out = {}
    for r in range(SIGNALS_FIRST, SIGNALS_FIRST + SIGNALS_ROWS):
        sym = ws.cell(row=r, column=1).value
        if sym and not isinstance(sym, str) or (isinstance(sym, str) and sym.strip()):
            out[str(sym).strip()] = (
                ws.cell(row=r, column=5).value,    # آخرین قیمت
                None,                              # قیمت پایانی (در Signals نیست)
                ws.cell(row=r, column=20).value,   # امتیاز کل
                ws.cell(row=r, column=21).value,   # سیگنال
                ws.cell(row=r, column=22).value,   # قدرت
            )
    # قیمت پایانی از Calculations
    if "Calculations" in wb.sheetnames:
        cs = wb["Calculations"]
        for rr in range(5, 5 + SIGNALS_ROWS):
            sym = cs.cell(row=rr, column=1).value
            if sym and str(sym).strip() in out:
                v = list(out[str(sym).strip()])
                v[1] = cs.cell(row=rr, column=5).value
                out[str(sym).strip()] = tuple(v)
    wb.close()
    return out


def link_options(stocks_path, options_path):
    data = _read_stocks(stocks_path)
    if not data:
        print("⚠ هیچ سیگنالی در فایل سهام یافت نشد. "
              "آیا فایل را بعد از ساخت در اکسل باز و ذخیره کرده‌اید؟ "
              "openpyxl مقدار محاسبه‌شده فرمول‌ها را ذخیره نمی‌کند.")
    wb = load_workbook(options_path)
    if "Underlying" not in wb.sheetnames:
        raise SystemExit("شیت Underlying در %s نیست." % options_path)
    ws = wb["Underlying"]
    today = datetime.date.today()
    written, missing = 0, []
    for r in range(UNDER_FIRST, UNDER_FIRST + UNDER_ROWS):
        sym = ws.cell(row=r, column=1).value
        if sym:
            key = str(sym).strip()
            if key in data:
                last, close, total, sig, strength = data[key]
                ws.cell(row=r, column=2, value=last)
                ws.cell(row=r, column=3, value=close)
                ws.cell(row=r, column=4, value=total)
                ws.cell(row=r, column=5, value=sig)
                ws.cell(row=r, column=6, value=strength)
                ws.cell(row=r, column=7, value=today).number_format = "yyyy-mm-dd"
                written += 1
            else:
                missing.append(key)
    wb.save(options_path)
    print("✓ %d نماد پایه از %s به %s منتقل شد."
          % (written, os.path.basename(stocks_path), os.path.basename(options_path)))
    if missing:
        print("⚠ این نمادها در فایل سهام نبودند: %s" % "، ".join(missing))
        print("  آن‌ها را به Watchlist فایل سهام اضافه کنید یا از شیت Underlying حذفشان کنید.")
    return written


def main(argv=None):
    ap = argparse.ArgumentParser(description="پل داده‌ای بین فایل‌های اکسل")
    ap.add_argument("--stocks", default="Stocks_Signals.xlsx")
    ap.add_argument("--options", default="Options_Signals.xlsx")
    args = ap.parse_args(argv)
    for p in (args.stocks, args.options):
        if not os.path.exists(p):
            raise SystemExit("فایل یافت نشد: %s" % p)
    link_options(args.stocks, args.options)
    print("\nتوجه: پس از این کار فایل آپشن را در اکسل باز کنید تا فرمول‌ها "
          "دوباره محاسبه شوند.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
